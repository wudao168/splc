import io
import json
from html import escape
from pathlib import Path
from .domain import row, rows


def export_stamp(content):
    """Keep stamp ink translucent and remove its paper background for overlays."""
    from PIL import Image
    with Image.open(io.BytesIO(content)) as source:
        image = source.convert('RGBA')
    pixels = [(r, g, b, round(a * (255 - min(r, g, b)) / 255 * .8))
              for r, g, b, a in image.getdata()]
    image.putdata(pixels)
    out = io.BytesIO()
    image.save(out, format='PNG')
    return out.getvalue()


def delivery_data(con, did, folder=None):
    document = row(con, 'SELECT * FROM deliveries WHERE id=?', (did,))
    lines = rows(con, 'SELECT * FROM delivery_lines WHERE delivery_id=? ORDER BY id', (did,))
    document['sender'] = dict(con.execute('SELECT name,phone,address FROM delivery_sender WHERE id=1').fetchone())
    document['lines'] = []
    for line in lines:
        snapshot = json.loads(line['snapshot'])
        projects = con.execute('SELECT project_code,subproject_code FROM order_lines WHERE id=?', (line['order_line_id'],)).fetchone()
        for key in ('project_code', 'subproject_code'):
            if not snapshot.get(key):
                snapshot[key] = projects[key] if projects else ''
        document['lines'].append({**snapshot, 'quantity': line['quantity']})
    document['po'] = '、'.join(dict.fromkeys(x['po'] for x in document['lines'] if x.get('po')))
    from .db import DATA
    stamp = con.execute("SELECT attachment_id FROM company_stamps WHERE kind='delivery'").fetchone()
    document['delivery_stamp'] = (Path(folder or DATA) / 'attachments' / stamp[0]).read_bytes() if stamp else None
    return document


def values(d):
    return [[i, x['name'], x['spec'], int(x['quantity']) if float(x['quantity']).is_integer() else x['quantity'], x['unit'],
             ' '.join(v for v in (x.get('carrier', ''), x.get('tracking', '')) if v), x.get('project_code', ''), x.get('subproject_code', ''), x.get('customer_code', '')]
            for i, x in enumerate(d['lines'], 1)]


def delivery_date(d):
    date = (d.get('shipped_at') or d['created_at'])[:10]
    year, month, day = date.split('-')
    return f'{year}年{int(month)}月{int(day)}日'


HEADERS = ['序号', '货物名称', '规格型号', '数量', '单位', '物流单号', '项目号', '子项目号', '备注']


def quote_data(con, oid):
    from decimal import Decimal, ROUND_HALF_UP
    order = row(con, 'SELECT * FROM orders WHERE id=?', (oid,))
    order['company'] = dict(con.execute('SELECT * FROM company_profile WHERE id=1').fetchone() or
                            {'name': '', 'address': '', 'phone': ''})
    order['lines'] = rows(con, 'SELECT * FROM order_lines WHERE order_id=? ORDER BY id', (oid,))
    for line in order['lines']:
        line['amount_cents'] = int((Decimal(str(line['quantity'])) * line['price_cents']).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    order['total_cents'] = sum(line['amount_cents'] for line in order['lines'])
    from .db import DATA
    from .domain import today
    stamp = con.execute("SELECT attachment_id FROM company_stamps WHERE kind='contract'").fetchone()
    order['contract_stamp'] = (DATA / 'attachments' / stamp[0]).read_bytes() if stamp else None
    order['quote_date'] = today()
    return order


def quote_rows(d):
    return [[i, l['name'], l['spec'], l['brand'], l['quantity'], l['unit'],
             l['price_cents'] / 100, l['amount_cents'] / 100,
             '\n'.join(v for v in (l.get('project_code', ''), l.get('subproject_code', '')) if v),
             l.get('remark', '')] for i, l in enumerate(d['lines'], 1)]


QUOTE_HEADERS = ['序号', '料品名称', '规格型号', '品牌', '数量', '单位', '单价（元）', '金额（元）', '项目号', '备注']


def quote_project_size(text):
    # The project column is 18 Excel units wide and the detail row is 48pt.
    from PIL import ImageFont
    font_path = Path('C:/Windows/Fonts/msyh.ttc')
    font = ImageFont.truetype(str(font_path), 40) if font_path.exists() else None
    for half_points in range(20, 1, -1):
        size = half_points / 2
        lines, width = 1, 0
        for char in str(text):
            if char == '\n':
                lines += 1
                width = 0
                continue
            advance = font.getlength(char) / 40 * size if font else size * (1 if ord(char) > 127 else .6)
            if width + advance > 88:
                lines += 1
                width = 0
            width += advance
        if lines * size * 1.5 <= 42:
            return size
    return 1


def xlsx_quote(d, delivery=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = '送货单' if delivery else '报价单'
    company = d['company']
    columns = 9 if delivery else 10
    amount_column = 9 if delivery else 8
    company_details = ' · '.join(filter(None, [company.get('name_en', ''), company['address'], company['phone']]))
    heading = f"{d['number']} · 送货单{'（已作废）' if d['status'] == 'void' else ''} · 客户 PO：{d['po']}" if delivery else f"{d['po']} · 报价单 · 第 {d['version']} 版"
    texts = [heading, company['name'], company_details,
             f"客户：{d['customer']}", f"收货地址：{d.get('address', '')}"]
    for i, text in enumerate(texts, 1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5 if i == 4 else columns)
        c = ws.cell(i, 1, text)
        c.data_type = 's'
        c.font = Font(name='Microsoft YaHei', size=14 if i == 2 else 10, bold=i == 2, color='243B55')
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.row_dimensions[i].height = 32
        if i in (2, 3):
            for column in range(1, columns + 1):
                ws.cell(i, column).fill = PatternFill('solid', fgColor='F4F8F9')
    ws.merge_cells(start_row=4, start_column=6, end_row=4, end_column=columns)
    ws['F4'] = f"联系人：{d['contact']} {d['phone']}"
    ws['F4'].data_type = 's'
    ws['F4'].font = Font(name='Microsoft YaHei', size=10, color='243B55')
    ws['F4'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.append(HEADERS if delivery else QUOTE_HEADERS)
    from openpyxl.worksheet.pagebreak import Break
    stamp_rows = []
    for index, record in enumerate(values(d) if delivery else quote_rows(d), 1):
        ws.append(record)
        if index < len(d['lines']) and (index % 12 == 0 or index == len(d['lines']) - 8) and len(d['lines']) - index >= 8:
            stamp_rows.append(ws.max_row - 2)
            ws.row_breaks.append(Break(id=ws.max_row))
    for cells in ws.iter_rows(min_row=6):
        for c in cells:
            if isinstance(c.value, str):
                c.data_type = 's'
            c.font = Font(name='Microsoft YaHei', size=10, bold=c.row == 6)
            if not delivery and c.row > 6 and c.column in (9, 10) and c.value:
                c.font = Font(name='Microsoft YaHei', size=quote_project_size(c.value))
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = Border(*([Side(style='thin', color='D7DEE5')] * 4))
            if c.row == 6:
                c.fill = PatternFill('solid', fgColor='E7F3F1')
            elif not delivery and c.column in (7, 8):
                c.number_format = '0.00'
        ws.row_dimensions[cells[0].row].height = 48
    ws.append(['数量合计' if delivery else '报价合计（含税）'] + [''] * (columns - 1))
    ws.cell(ws.max_row, amount_column, sum(l['quantity'] for l in d['lines']) if delivery else d['total_cents'] / 100)
    ws.cell(ws.max_row, amount_column).number_format = '0.######' if delivery else '0.00'
    ws.cell(ws.max_row, amount_column).alignment = Alignment(horizontal='center')
    for cells in ws.iter_rows(min_row=7, max_row=ws.max_row, min_col=8 if delivery else 7, max_col=9 if delivery else 8):
        for c in cells:
            c.alignment = Alignment(horizontal='center', vertical='center')
    if d['note']:
        ws.append(['备注：' + d['note']])
        ws.cell(ws.max_row, 1).data_type = 's'
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=columns)
    from openpyxl.utils import get_column_letter
    for i, width in enumerate([5, 16, 19, 7, 5, 18, 9, 9, 10] if delivery else [8, 16, 19, 8, 10, 7, 9, 10, 18, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    footer = ws.max_row + 2
    fields = [(1, 3, '发货人：' + (d.get('sender', {}).get('name') or '________________')), (4, 6, '客户签字：________________'), (7, 9, '日期：' + delivery_date(d))] if delivery else [(1, 5, '报价人：' + d.get('quoted_by', '________________')), (6, 10, '日期：' + d.get('quote_date', ''))]
    ws.row_dimensions[footer].height = 120
    for start, end, label in fields:
        ws.merge_cells(start_row=footer, start_column=start, end_row=footer, end_column=end)
        cell = ws.cell(footer, start, label)
        cell.font = Font(name='Microsoft YaHei', size=11)
        cell.alignment = Alignment(vertical='center')
    if d.get('contract_stamp'):
        from openpyxl.drawing.image import Image
        for stamp_row in [*stamp_rows, footer]:
            image = Image(io.BytesIO(export_stamp(d['contract_stamp'])))
            ratio = min(153 / image.width, 153 / image.height)
            image.width *= ratio
            image.height *= ratio
            ws.add_image(image, f'F{stamp_row}')
    ws.freeze_panes = 'A7'
    ws.print_title_rows = '1:6'
    ws.print_area = f'A1:{"I" if delivery else "J"}{ws.max_row}'
    ws.page_setup.orientation = 'portrait'
    from openpyxl.worksheet.page import PageMargins
    ws.page_margins = PageMargins(left=1/2.54, right=1/2.54, top=1/2.54, bottom=1/2.54, header=0, footer=0)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = '第 &P 页 / 共 &N 页'
    ws.oddFooter.center.size = 9
    ws.page_margins.footer = .15
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pdf_quote(d, delivery=False):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, KeepInFrame
    from reportlab.lib import colors
    font_path = Path('C:/Windows/Fonts/msyh.ttc')
    if font_path.exists():
        if 'CaidanChinese' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('CaidanChinese', str(font_path), subfontIndex=0))
        font_name = 'CaidanChinese'
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        font_name = 'STSong-Light'
    style = ParagraphStyle('quote', fontName=font_name, fontSize=9, leading=13, wordWrap='CJK')
    title = ParagraphStyle('quote-title', parent=style, fontSize=18, leading=26, alignment=1)
    p = lambda text: Paragraph(escape(str(text)).replace('\n', '<br/>'), style)
    company = d['company']
    company_details = ' · '.join(filter(None, [company.get('name_en', ''), company['address'], company['phone']]))
    company_heading = ParagraphStyle('company-heading', parent=style, fontSize=12, leading=18)
    company_block = Table([[Paragraph(escape(company['name']), company_heading)], [p(company_details)]], colWidths=[538.58])
    company_block.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F4F8F9')),
                                      ('LINEBEFORE', (0, 0), (0, -1), 2, colors.HexColor('#008B83')),
                                      ('LEFTPADDING', (0, 0), (-1, -1), 12),
                                      ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
    customer_block = Table([[p(f"客户：{d['customer']}"), p(f"联系人：{d['contact']} {d['phone']}")],
                            [p(f"收货地址：{d.get('address', '')}"), '']], colWidths=[269.29, 269.29])
    customer_block.setStyle(TableStyle([('SPAN', (0, 1), (1, 1)), ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                       ('LEFTPADDING', (0, 0), (-1, -1), 0),
                                       ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    heading = f"{d['number']} · 送货单{'（已作废）' if d['status'] == 'void' else ''} · 客户 PO：{d['po']}" if delivery else f"{d['po']} · 报价单 · 第 {d['version']} 版"
    po_heading = Table([[p(heading)]], colWidths=[538.58])
    po_heading.setStyle(TableStyle([('LEFTPADDING', (0, 0), (-1, -1), 0)]))
    story = [po_heading, Spacer(1, 12), company_block,
             Spacer(1, 10), customer_block, Spacer(1, 12)]
    cell_style = ParagraphStyle('quote-cell', parent=style, alignment=1)
    records = [[Paragraph(escape(v), cell_style) for v in (HEADERS if delivery else QUOTE_HEADERS)]]
    records.extend([[p(f'{v:g}' if i == 3 else v) for i, v in enumerate(r)] for r in values(d)] if delivery else [[p(f'{v:.2f}' if i in (6, 7) else f'{v:g}' if i == 4 else v) for i, v in enumerate(r)] for r in quote_rows(d)])
    right = ParagraphStyle('quote-amount', parent=style, alignment=1)
    for record in records[1:]:
        for i in range(len(record)):
            record[i].style = right
    total = [p('数量合计' if delivery else '报价合计（含税）')] + [''] * (8 if delivery else 9)
    total[8 if delivery else 7] = Paragraph(f"{sum(l['quantity'] for l in d['lines']):g}" if delivery else f"{d['total_cents'] / 100:.2f}", right)
    records.append(total)
    widths = [25, 85, 90, 35, 28, 105, 55, 60, 55.58] if delivery else [35, 75, 85, 35, 40, 30, 50, 55, 75, 58.58]
    for record in records[1:-1]:
        for i, value in enumerate(record):
            record[i] = KeepInFrame(widths[i] - 12, 30, [value], mode='shrink')
    table = Table(records, colWidths=widths, rowHeights=[32] + [42] * len(d['lines']) + [32], repeatRows=1)
    table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#D7DEE5')),
                              ('SPAN', (0, -1), (7 if delivery else 6, -1)),
                              ('NOSPLIT', (0, -2), (-1, -1)),
                              ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E7F3F1')),
                              ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                              ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
    story.append(table)
    if d['note']:
        story.append(p('备注：' + d['note']))
    from reportlab.platypus import Flowable
    from reportlab.lib.utils import ImageReader
    class QuoteSignature(Flowable):
        def __init__(self):
            super().__init__()
            self.width, self.height = 538.58, 130

        def draw(self):
            self.canv.setFont(font_name, 11)
            self.canv.drawString(0, 53, '发货人：' + (d.get('sender', {}).get('name') or '________________') if delivery else '报价人：' + d.get('quoted_by', '________________'))
            if delivery:
                self.canv.drawString(170, 53, '客户签字：________________')
            self.canv.drawString(340, 53, '日期：' + (delivery_date(d) if delivery else d.get('quote_date', '')))
            if d.get('contract_stamp'):
                self.canv.drawImage(ImageReader(io.BytesIO(export_stamp(d['contract_stamp']))), 365, 5,
                                    width=113.39, height=113.39, preserveAspectRatio=True, mask='auto')
                self.canv.drawString(340, 53, '日期：' + (delivery_date(d) if delivery else d.get('quote_date', '')))
    story.extend([Spacer(1, 12), QuoteSignature()])
    out = io.BytesIO()
    from reportlab.pdfgen.canvas import Canvas
    class NumberedCanvas(Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.pages = []

        def showPage(self):
            self.pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            count = len(self.pages)
            for state in self.pages:
                self.__dict__.update(state)
                self.setFont(font_name, 9)
                self.drawCentredString(A4[0] / 2, 14, f'第 {self._pageNumber} 页 / 共 {count} 页')
                if self._pageNumber < count and d.get('contract_stamp'):
                    self.drawImage(ImageReader(io.BytesIO(export_stamp(d['contract_stamp']))),
                                   A4[0] - 28.3465 - 113.39, 32, width=113.39, height=113.39,
                                   preserveAspectRatio=True, mask='auto')
                super().showPage()
            super().save()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=28.3465, rightMargin=28.3465, topMargin=28.3465,
                      bottomMargin=28.3465).build(story, canvasmaker=NumberedCanvas)
    return out.getvalue()

BILLING_FIELDS = [('name', '公司名称'), ('bank_name', '开户行'), ('bank_account', '账号'),
                  ('tax_number', '税号'), ('address', '地址'), ('phone', '电话'), ('email', '邮箱')]
CUSTOMER_BILLING_FIELDS = [('title', '发票抬头'), ('tax_number', '纳税人识别号'), ('address', '注册地址'),
                           ('phone', '注册电话'), ('bank_name', '开户银行'), ('bank_account', '银行账号'), ('email', '收票邮箱')]


def xlsx_billing(profile, fields=BILLING_FIELDS, title='开票付款资料'):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:B1')
    ws['A1'] = title
    ws['A1'].font = Font(name='Microsoft YaHei', size=20, bold=True, color='183451')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 52
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 66
    for i, (key, label) in enumerate(fields, 3):
        for column, value in enumerate((label, profile.get(key, '')), 1):
            cell = ws.cell(i, column, value)
            cell.data_type = 's'
            cell.number_format = '@'
            cell.font = Font(name='Microsoft YaHei', size=12, bold=column == 1, color='183451')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if i % 2:
                cell.fill = PatternFill('solid', fgColor='F1F5F8')
        value = profile.get(key, '')
        visual_length = sum(2 if ord(char) > 127 else 1 for char in value)
        ws.row_dimensions[i].height = max(42, ((visual_length + 55) // 56) * 22 + 16)
    ws.print_options.horizontalCentered = True
    ws.print_area = 'A1:B9'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pdf_billing(profile, fields=BILLING_FIELDS, title='开票付款资料'):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    font_path = Path('C:/Windows/Fonts/msyh.ttc')
    if font_path.exists():
        if 'CaidanChinese' not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont('CaidanChinese', str(font_path), subfontIndex=0))
        font_name = 'CaidanChinese'
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        font_name = 'STSong-Light'
    body = ParagraphStyle('billing', fontName=font_name, fontSize=12, leading=21, wordWrap='CJK')
    title_style = ParagraphStyle('billing-title', parent=body, fontSize=22, leading=32, alignment=1)
    paragraph = lambda value: Paragraph(escape(value).replace('\n', '<br/>'), body)
    table = Table([[paragraph(label), paragraph(profile.get(key, ''))] for key, label in fields], colWidths=[110, 385])
    table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                              ('TOPPADDING', (0, 0), (-1, -1), 17),
                              ('BOTTOMPADDING', (0, 0), (-1, -1), 17),
                              ('LEFTPADDING', (0, 0), (-1, -1), 12),
                              ('RIGHTPADDING', (0, 0), (-1, -1), 12),
                              ('ROWBACKGROUNDS', (0, 0), (-1, -1), [colors.HexColor('#F1F5F8'), colors.white])]))
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=58, bottomMargin=50,
                            title=title, author=profile.get('name') or profile.get('title', ''))
    doc.build([Paragraph(title, title_style), Spacer(1, 30), table])
    return out.getvalue()


def delivery_template(d):
    return {**d, 'company': {'name': d['company'], 'name_en': d['company_en'],
                            'address': d['company_address'], 'phone': d['company_phone']},
            'contract_stamp': d.get('delivery_stamp')}


def xlsx_delivery(d):
    return xlsx_quote(delivery_template(d), delivery=True)


def pdf_delivery(d):
    return pdf_quote(delivery_template(d), delivery=True)


def html_delivery(d):
    e = lambda s: escape(str(s))
    heads = ''.join(f'<th>{h}</th>' for h in HEADERS)
    body = ''.join('<tr>' + ''.join(f'<td>{e(v)}</td>' for v in r) + '</tr>' for r in values(d))
    return f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>{e(d['number'])}</title>
    <style>body{{font:14px "Microsoft YaHei",sans-serif;margin:32px;color:#172b4d}}h1,h2,.supplier{{text-align:center}}.supplier{{margin:4px}}.delivery-info{{display:grid;grid-template-columns:65.3% 34.7%;grid-auto-rows:32px;align-items:center}}table{{width:100%;table-layout:fixed;border-collapse:collapse}}th,td{{border:1px solid #ccd5df;padding:9px;word-break:break-word;text-align:center;vertical-align:middle;height:56px;box-sizing:border-box}}th{{background:#e7f3f1}}button{{padding:10px 20px;margin-bottom:20px}}@media print{{button{{display:none}}thead{{display:table-header-group}}tr{{break-inside:avoid}}}}@page{{size:A4 portrait;margin:12mm}}</style>
    <button onclick="window.print()">打印 / 另存 PDF</button><h1>{e(d['company'])}</h1><p class="supplier">{e(d['company_en'])}</p><p class="supplier">{e(d['company_address'])}　{'TEL：' + e(d['company_phone']) if d['company_phone'] else ''}</p><h2>送货单{'（已作废）' if d['status']=='void' else ''}</h2>
    <div class="delivery-info"><span>客户名称：{e(d['customer'])}</span><span></span><span>联系人：{e(d['contact'])} {e(d['phone'])}</span><span></span><span>收货地址：{e(d['address'])}</span><span>客户PO：{e(d['po'])}</span><span>送货单号：{e(d['number'])}</span><span>日期：{e(d['created_at'][:10])}</span></div>
    <table><colgroup><col style="width:4.6%"><col style="width:15.8%"><col style="width:16.7%"><col style="width:6.5%"><col style="width:5.2%"><col style="width:19.5%"><col style="width:10.2%"><col style="width:11.1%"><col style="width:10.4%"></colgroup><thead><tr>{heads}</tr></thead><tbody>{body}</tbody></table><p>{e(d['note'])}</p><p>发货人：{e(d.get('sender', {}).get('name', ''))}</p></html>'''.encode()

