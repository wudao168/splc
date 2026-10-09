"""Dashboard reports, shared by the screen and Excel export."""
import io
from . import domain as dm

TITLES = {'orders':'客户订单汇总','pending':'待采购明细','purchases':'采购明细','invoices':'发票汇总','delivery':'交付明细'}

def build(con, params):
    kind = params.get('kind', 'orders')
    dm.require(kind in TITLES, '报表类型无效')
    start, end = params.get('start', ''), params.get('end', '')
    dm.valid_date(start); dm.valid_date(end)
    dm.require(not start or not end or start <= end, '开始日期不能晚于结束日期')
    data = dm.get_state(con)
    orders = {o['id']:o for o in data['orders']}
    lines = {l['id']:l for l in data['order_lines']}
    result = []
    def matches(date, related):
        wanted = (params.get('customer') or '').strip().lower()
        return (not start or date >= start) and (not end or date <= end) and any((not wanted or wanted in o['customer'].lower()) and (not params.get('po') or params['po'].lower() in o['po'].lower()) for o in related)
    if kind in ('orders','pending','delivery'):
        headers = {'orders':['客户','客户 PO','订单日期','订单状态','报价金额（元）','采购成本（元）','毛利（元）','毛利率','待采购料品项数','实际发货数量 / 需求数量'], 'pending':['客户','客户 PO','订单日期','料品','规格','单位','需求数量','已采购数量','待采购数量'], 'delivery':['客户','客户 PO','订单日期','料品','规格','单位','需求数量','已开单数量','实际发货数量','待交付数量']}[kind]
        for order in orders.values():
            if order['status'] != 'confirmed' or order.get('archived_at'): continue
            date = order['created_at'][:10]
            if not matches(date, [order]): continue
            items = [l for l in lines.values() if l['order_id'] == order['id']]
            prefix = [order['customer'], order['po'], date]
            if kind == 'orders':
                quote = sum(round(l['demand_quantity'] * l['price_cents']) for l in items)
                cost = sum(l['cost_cents'] for l in items)
                pending_count = sum(l['demand_quantity'] > l['purchased'] + 1e-6 for l in items)
                complete = bool(items) and pending_count == 0
                profit = (quote-cost)/100 if complete else None
                rate = f'{(quote-cost)/quote*100:.2f}%' if complete and quote else None
                result.append(prefix+['已确认',quote/100,cost/100,profit,rate,pending_count, f"{sum(l['dispatched_quantity'] for l in items):g} / {sum(l['demand_quantity'] for l in items):g}"])
            else:
                for line in items:
                    demand = line['demand_quantity']
                    if kind == 'pending':
                        pending = max(0, demand-line['purchased'])
                        if pending <= 1e-6: continue
                        values = [demand,line['purchased'],round(pending,6)]
                    else:
                        active = {d['id'] for d in data['deliveries'] if d['status']=='active'}
                        opened = sum(d['quantity'] for d in data['delivery_lines'] if d['order_line_id']==line['id'] and d['delivery_id'] in active and not d.get('replacement_case_id'))
                        values = [demand,opened,line['dispatched_quantity'],max(0,demand-line['dispatched_quantity'])]
                    result.append(prefix+[line['name'],line['spec'],line['unit']]+values)
    else:
        headers = ['平台','店铺','平台订单号','关联客户','关联 PO','采购日期','实付款（元）'] + (['有效采购数量','退货数量','物流状态','料品明细（数量 @ 单价）'] if kind=='purchases' else ['已收票金额（元）','待收票金额（元）','开票状态','收票状态'])
        items_by_id = {i['id']: i for i in data['items']}
        for p in data['purchases']:
            pls = [l for l in data['purchase_lines'] if l['purchase_id']==p['id']]
            related = list({lines[l['order_line_id']]['order_id']:orders[lines[l['order_line_id']]['order_id']] for l in pls if l['order_line_id'] in lines}.values())
            if not matches(p['purchased_date'], related or [{'customer':'','po':''}]): continue
            packages = {k['tracking']:k for k in (p.get('taobao_source') or {}).get('packages', [])}
            lids = {l['id'] for l in pls}
            package_ids = {l['package_id'] for l in data['package_lines'] if l['purchase_line_id'] in lids}
            packages.update({k['tracking']:k for k in data['packages'] if k['id'] in package_ids})
            transit = any(k.get('status') in ('待揽收','运输中','派送中') for k in packages.values())
            if params.get('scope')=='transit' and not transit: continue
            if params.get('scope')=='missing' and not (p['remaining_cents']>0 and p['invoice_stage']!='不需开票'): continue
            prefix = [p['platform'],p['shop'],p['platform_order'],'、'.join(dict.fromkeys(o['customer'] for o in related)),'、'.join(o['po'] for o in related),p['purchased_date'],p['amount_cents']/100]
            if kind=='purchases':
                returned = sum(c['quantity'] for c in data['purchase_cases'] if c['purchase_line_id'] in lids and c['kind']=='supplier_return' and c['status']=='completed')
                detail = '；'.join(
                    f"{lines[l['order_line_id']]['name'] if l['order_line_id'] in lines else items_by_id.get(l['item_id'], {}).get('name', l.get('item_name') or '备货料品')} "
                    f"{l['quantity']:g} @ {((l.get('unit_price_cents') if l.get('unit_price_cents') is not None else (round(l['cost_cents'] / l['quantity']) if l['quantity'] else l['cost_cents'])) / 100):.2f}"
                    for l in pls) or '暂无明细'
                values = [sum(l['quantity'] for l in pls),returned,'；'.join(f"{k.get('carrier','')} {k['tracking']} {k.get('status','')}" for k in packages.values()) or '暂无运单',detail]
            else:
                values = [p['received_cents']/100,p['remaining_cents']/100,p['invoice_stage'],p['receipt_status']]
            result.append(prefix+values)
    date_index = 2 if kind in ('orders','pending','delivery') else 5
    result.sort(key=lambda r:r[date_index], reverse=True)
    summary = None
    if kind == 'orders':
        completed = [r for r in result if r[6] is not None]
        quote, cost = [round(sum(r[i] for r in completed), 2) for i in (4,5)]
        shipped, demand = [sum(float(r[9].split(' / ')[i]) for r in completed) for i in (0,1)]
        summary = ['汇总（采购已完成）',None,None,None,quote,cost,round(quote-cost,2),f'{(quote-cost)/quote*100:.2f}%' if quote else None,0,f'{shipped:g} / {demand:g}']
    elif kind == 'invoices':
        summary = ['汇总',None,None,None,None,None]+[round(sum(r[i] for r in result),2) for i in (6,7,8)]+[None,None]
    elif kind == 'pending':
        # 与客户订单汇总一致：数量列直接相加（单位不同的明细同样累加，便于快速看总量）。
        summary = [f'汇总（{len(result)} 项）',None,None,None,None,None]+[round(sum(r[i] for r in result),6) for i in (6,7,8)]
    elif kind == 'delivery':
        summary = [f'汇总（{len(result)} 项）',None,None,None,None,None]+[round(sum(r[i] for r in result),6) for i in (6,7,8,9)]
    elif kind == 'purchases':
        summary = [f'汇总（{len(result)} 笔）',None,None,None,None,None,round(sum(r[6] for r in result),2),round(sum(r[7] for r in result),6),round(sum(r[8] for r in result),6),None,None]
    return {'summary':summary, 'title':TITLES[kind], 'headers':headers, 'rows':result, 'date_label':'订单创建日期' if date_index==2 else '采购日期', 'customers':sorted({o['customer'] for o in orders.values()})}

def xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    book = Workbook(); sheet = book.active; sheet.title = report['title']
    sheet.append(report['headers'])
    for row in ([report['summary']] if report.get('summary') else []) + report['rows']:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
    for cell in sheet[1]: cell.font = Font(bold=True)
    for row in sheet:
        for cell in row: cell.alignment = Alignment(vertical='center', wrap_text=True)
    for index in range(1,len(report['headers'])+1): sheet.column_dimensions[get_column_letter(index)].width=24
    sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
    out=io.BytesIO(); book.save(out); return out.getvalue()
