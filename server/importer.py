import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

FIELDS = {
    'name': ['料品名称', '物料名称', '货物名称', '商品名称', '产品名称', '品名', '名称'],
    'spec': ['料品规格', '规格型号', '产品型号', '型号规格', '规格', '型号', '型号或图号', '规格/型号', '型号/规格'],
    'brand': ['品牌', '厂牌', '制造商', '品牌或材质'],
    'description': ['料品描述', '物料描述', '商品描述', '描述', '说明', '技术要求', '技术说明'],
    'quantity': ['采购数量', '需求数量', '订购数量', '数量'],
    'unit': ['计量单位', '单位'],
    'customer_code': ['客户料号', '物料编码', '物料编号', '料号', '物料代码'],
    'project_code': ['项目编码', '项目号', '项目编号'],
    'subproject_code': ['子项目号', '子项目', '子项目编码', '子项目编号'],
    'remark': ['备注', '料品备注', '报价备注'],
    'price': ['报价单价', '含税单价', '单价', '单价（含税）', '单价(含税)'],
}
PO_HEADERS = {'客户po', 'po', 'po单号', '客户po/询价单号', '采购单号', '采购订单号', '订单编号', '订单号', '单据编号', '询价单号', '客户订单号'}


def normalize_header(value):
    return re.sub(r'[\s：:.]+', '', clean(value)).lower()


def clean(value):
    return str(value if value is not None else '').strip()


def coordinate_table(page, title):
    words = page.extract_words()
    for word in words:
        if not guess_mapping([word['text']]).get('name') == 0:
            continue
        headers = sorted([w for w in words if abs(w['top'] - word['top']) < 3], key=lambda w: w['x0'])
        sequence = next((w for w in words if w['text'] == '序' and abs(w['top'] - word['top']) <= 6 and w['x1'] < headers[0]['x0']), None)
        if sequence:
            headers.insert(0, {**sequence, 'text': '序号'})
        mapping = guess_mapping([w['text'] for w in headers])
        if 'quantity' not in mapping:
            continue
        centers = [(w['x0'] + w['x1']) / 2 for w in headers]
        def column(w):
            x = (w['x0'] + w['x1']) / 2
            return sum(x > (a + b) / 2 for a, b in zip(centers, centers[1:]))
        bottom = min((w['top'] for w in words if w['top'] > word['bottom'] and re.search('合计|总计', w['text'])), default=page.height)
        body = [w for w in words if w['top'] > word['bottom'] and w['top'] < bottom]
        anchors = sorted([w for w in body if column(w) == mapping['quantity'] and re.fullmatch(r'\d+(?:\.\d+)?', w['text'])], key=lambda w: w['top'])
        if not anchors:
            continue
        rows = []
        for i, anchor in enumerate(anchors):
            low = (anchors[i-1]['top'] + anchor['top']) / 2 if i else word['bottom']
            high = (anchor['top'] + anchors[i+1]['top']) / 2 if i+1 < len(anchors) else bottom
            cells = [[] for _ in headers]
            for w in sorted(body, key=lambda w: (w['top'], w['x0'])):
                if low <= w['top'] < high:
                    cells[column(w)].append(w['text'])
            rows.append([ ''.join(cell) for cell in cells ])
        return table_from_grid([[w['text'] for w in headers], *rows], title)


def guess_mapping(headers):
    mapping = {}
    for i, title in enumerate(headers):
        title = re.sub(r'\s+', '', clean(title))
        for field, aliases in FIELDS.items():
            if title in aliases and field not in mapping:
                mapping[field] = i
    return mapping


def table_from_grid(grid, title):
    candidates = []
    for i, headers in enumerate(grid[:80]):
        mapping = {}
        for col, header in enumerate(headers):
            field = next(iter(guess_mapping([header])), None)
            if field:
                if field in mapping:
                    candidates.append((i, mapping))
                    mapping = {}
                mapping[field] = col
        candidates.append((i, mapping))
    def score(candidate):
        i, mapping = candidate
        if 'name' not in mapping or 'quantity' not in mapping:
            return (0, 0)
        populated = sum(bool(clean(r[mapping['name']])) and bool(clean(r[mapping['quantity']]))
                        for r in grid[i + 1:] if len(r) > max(mapping['name'], mapping['quantity']))
        return (len(mapping), populated)
    best = max(candidates, key=score, default=(0, {}))
    idx, mapping = best
    if 'name' not in mapping or 'quantity' not in mapping:
        return None
    data = []
    for i, raw in enumerate(grid[idx + 1:], idx + 2):
        values = [clean(x) for x in raw]
        if not any(values):
            continue
        if 'quantity' in guess_mapping(values):
            continue
        name = values[mapping['name']] if len(values) > mapping['name'] else ''
        quantity = values[mapping['quantity']] if len(values) > mapping['quantity'] else ''
        if not name or re.search(r'^(合计|总计|小计|备注|出货人|客户[（(]?签)', name):
            continue
        if not quantity and len([v for v in values if v]) <= 1:
            continue
        data.append({'row': i, 'values': values})
    po_columns = [i for i, h in enumerate(grid[idx]) if normalize_header(h) in PO_HEADERS]
    candidates = list(dict.fromkeys(r['values'][i] for r in data for i in po_columns if len(r['values']) > i and r['values'][i]))
    return {'title': title, 'header_row': idx + 1, 'headers': [clean(x) for x in grid[idx]], 'mapping': mapping, 'rows': data,
            'po_candidates': candidates, **({'po_column': po_columns[0]} if po_columns else {}), 'metadata': {'po': candidates[0]} if len(candidates) == 1 else {}}


def metadata(text):
    result = {}
    for field, expression in {
        'customer': r'(?:客户名称|客户|购货单位)[：:]\s*([^\n\t]+)',
        'po': r'(?:客户\s*PO|PO\s*单号|采购订单号|采购单号|客户订单号|订单编号|订单号|单据编号|询价单号|\bP\.?O\.?)[：:]?[ \t]*([A-Za-z0-9_-]{5,})',
        'address': r'(?:收货地址|地址)[：:]\s*([^\n\t]+)',
        'contact': r'(?:联系人|收货人)[：:]\s*([^\n\t]+)',
    }.items():
        match = re.search(expression, text, re.I)
        if match:
            result[field] = match.group(1).strip()
    if 'customer' not in result:
        match = re.search(r'需\s*方[：:]\s*(.+?)(?=\s{2,}供\s*方|$)', text, re.M)
        if match:
            result['customer'] = match.group(1).strip()
    if result.get('contact'):
        match = re.search(r'1\d{10}', result['contact'])
        if match:
            result['phone'] = match.group()
            result['contact'] = result['contact'].replace(match.group(), '').strip()
    return result


def parse_file(content, filename):
    suffix = Path(filename).suffix.lower()
    tables, pieces, warnings = [], [], []
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 400 * 1024 * 1024:
                raise ValueError('Excel 解压后过大，请拆分后导入')
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        try:
            for sheet in wb:
                max_row, max_column = content_bounds(content, sheet)
                if max_row > 10000 or max_column > 150:
                    warnings.append(f'{sheet.title} 有效内容范围超过 10000 行或 150 列，请精简后导入')
                    continue
                if not max_row:
                    continue
                sheet.reset_dimensions()
                grid = list(sheet.iter_rows(max_row=max_row, max_col=max_column, values_only=True))
                pieces.append('\n'.join('\t'.join(clean(x) for x in r) for r in grid))
                table = table_from_grid(grid, sheet.title)
                if table:
                    tables.append(table)
        finally:
            wb.close()
    elif suffix == '.pdf':
        import pdfplumber
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('请先解除 PDF 密码保护后再导入')
        if len(reader.pages) > 100:
            raise ValueError('PDF 超过 100 页，请拆分导入')
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for n, page in enumerate(reader.pages, 1):
                text = page.extract_text(extraction_mode='layout') or ''
                pieces.append(text)
                positioned = coordinate_table(pdf.pages[n - 1], f'第 {n} 页')
                if positioned:
                    tables.append(positioned)
                    continue
                grids = pdf.pages[n - 1].extract_tables()
                grids.append([re.split(r'\s{2,}|\t', s.strip()) for s in text.splitlines() if s.strip()])
                for grid in grids:
                    table = table_from_grid(grid, f'第 {n} 页')
                    if table:
                        quantity_col = table['mapping']['quantity']
                        table['rows'] = [r for r in table['rows'] if len(r['values']) > quantity_col
                                         and re.fullmatch(r'\d+(?:\.\d+)?', r['values'][quantity_col])]
                        if table['rows']:
                            tables.append(table)
                            break
        if not any(x.strip() for x in pieces):
            warnings.append('此 PDF 没有可提取文字，可能为扫描件。OCR 尚未接入，请对照原件手动录入。')
        else:
            warnings.append('PDF 表格可能存在错列或跨页漏行，请逐项对照原件。')
    else:
        raise ValueError('客户文件目前支持 .xlsx 和 .pdf；旧版 .xls 请另存为 .xlsx')
    raw = '\n'.join(pieces)
    meta = metadata(raw)
    candidates = list(dict.fromkeys(p for t in tables for p in t['po_candidates']))
    if len(candidates) == 1:
        meta['po'] = candidates[0]
    elif len(candidates) > 1:
        meta.pop('po', None)
        warnings.append('文件中包含多个 PO，请按 PO 筛选对应明细后分别保存，不要合并成一张订单。')
    if not tables:
        warnings.append('未识别到同时包含名称和数量的表头，可粘贴制表符分隔的明细或手动添加料品。')
    warnings.append('品牌、规格及数量请核对；源文件缺少的字段不会自动补造。')
    return {'tables': tables, 'text': raw[:100000], 'metadata': meta, 'warnings': warnings}


def parse_paste(text, kind):
    if kind == 'items':
        grid = [re.split(r'\t| {2,}', s.strip()) for s in text.splitlines() if s.strip()]
        table = table_from_grid(grid, '粘贴明细')
        if not table:
            raise ValueError('未找到料品名称和数量表头。请从 Excel 连同表头复制，列之间使用制表符。')
        return {'tables': [table]}
    text = re.sub(r'(?<=[\u4e00-\u9fff])[ \t]+(?=[\u4e00-\u9fff])', '', text)
    result = {}
    patterns = {
        'platform_order': r'(?:订单编号|订单号)[：:\s]*([A-Za-z0-9_-]{5,})',
        'shop': r'(?:店铺名称|店铺名|店铺|商家名称|商家|卖家昵称|卖家)[：:]\s*([^\r\n]+)',
        'amount': r'(?:实付款|实际付款|实付金额)[^\d\r\n]{0,20}(?:\r?\n[ \t]*)?[￥¥]?[ \t]*([\d,]+(?:\.\d{1,2})?)',
        'purchased_date': r'(?:订单创建时间|创建时间|下单时间|付款时间)[：:\s]*(\d{4}[/-]\d{1,2}[/-]\d{1,2})',
        'tracking': r'(?:快递单号|运单号码|运单号|物流单号|物流编号)[：:\s]*([A-Za-z0-9-]{6,})',
        'carrier': r'(?:快递公司|物流公司|承运公司)[：:]\s*([^\r\n]+)',
    }
    for field, pattern in patterns.items():
        match = re.search(pattern, text)
        if match:
            value = match.group(1).strip()
            if field == 'purchased_date':
                year, month, day = re.split(r'[/-]', value)
                value = f'{year}-{int(month):02d}-{int(day):02d}'
            result[field] = value.replace(',', '') if field == 'amount' else value
    if 'shop' not in result:
        match = re.search(r'(?m)^展开全部商品[ \t]*\r?\n[ \t]*([^\r\n]{2,80})\r?\n[ \t]*查看交易快照', text)
        if match:
            result['shop'] = match.group(1).strip()
    if 'tracking' not in result:
        match = re.search(r'(?mi)^[ \t]*((?:中通|圆通|申通|韵达|顺丰|极兔|京东|德邦|邮政|EMS|菜鸟)[^\r\n]{0,10})\r?\n[ \t]*([A-Z0-9-]{8,30})[ \t]*$', text)
        if match:
            result['carrier'], result['tracking'] = match.group(1).strip(), match.group(2)
    if not result:
        raise ValueError('未识别到订单号、店铺、实付款或下单时间。请从订单详情页复制包含这些字段的文字。')
    return result


def content_bounds(content, sheet):
    """Ignore dimension metadata and cells that contain only formatting."""
    from openpyxl.utils.cell import coordinate_to_tuple
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    max_row = max_column = 0
    # The read-only worksheet stores its XML package path here.
    with zipfile.ZipFile(io.BytesIO(content)) as archive, archive.open(sheet._worksheet_path) as source:
        for _, element in ET.iterparse(source, events=('end',)):
            if element.tag == ns + 'c':
                has_value = any(node.text and node.text.strip() for node in element.iter()
                                if node.tag in (ns + 'v', ns + 't', ns + 'f'))
                if has_value:
                    row, column = coordinate_to_tuple(element.attrib['r'])
                    max_row, max_column = max(max_row, row), max(max_column, column)
                element.clear()
            elif element.tag == ns + 'row':
                element.clear()
    return max_row, max_column
