import io
import unittest
from openpyxl import load_workbook
from tests.test_purchase_associations import PurchaseAssociationTests
from server import reports, domain as dm

class ReportTests(PurchaseAssociationTests):
    def test_pending_cost_filters_and_export(self):
        pending = reports.build(self.con, {'kind':'pending'})
        self.assertEqual(len(pending['rows']),1)
        self.assertEqual(pending['rows'][0][-1],3)
        purchase = reports.build(self.con, {'kind':'purchases'})
        self.assertEqual(purchase['rows'][0][6],20)
        self.assertEqual(reports.build(self.con, {'kind':'orders','customer':'不存在'})['rows'],[])
        self.assertEqual(len(reports.build(self.con, {'kind':'orders','po':'EDIT'})['rows']),1)
        with self.assertRaises(ValueError): reports.build(self.con, {'start':'2026-10-07','end':'2026-10-06'})
        sheet = load_workbook(io.BytesIO(reports.xlsx(purchase))).active
        self.assertEqual(sheet.cell(2,7).value,20)
        self.assertEqual(sheet.cell(2,3).data_type,'s')

    def test_delivery_invoice_and_source_logistics(self):
        dm.save_taobao_source(self.con,self.pid,dict(platform_order='12345678901',packages=[dict(carrier='快递',tracking='TRACK123456',status='运输中')]))
        self.assertEqual(len(reports.build(self.con,{'kind':'purchases','scope':'transit'})['rows']),1)
        self.assertEqual(len(reports.build(self.con,{'kind':'invoices','scope':'missing'})['rows']),1)
        dm.create_deliveries(self.con,dict(company='公司',lines=[dict(order_line_id=self.old,quantity=1)]))
        row = next(r for r in reports.build(self.con,{'kind':'delivery'})['rows'] if r[3]=='旧料品')
        self.assertEqual(row[-3:],[1,1,1])

    def test_order_profit_requires_complete_procurement(self):
        row = reports.build(self.con, {'kind':'orders'})['rows'][0]
        self.assertEqual(row[3], '已确认')
        self.assertEqual(row[6:8], [None, None])
        self.con.execute('UPDATE order_lines SET price_cents=1000')
        dm.create_purchase(self.con, dict(platform='淘宝', platform_order='22345678901', shop='店铺', amount='10', lines=[dict(order_line_id=self.new, quantity=3, cost='10')]))
        report = reports.build(self.con, {'kind':'orders'})
        self.assertEqual(report['rows'][0][6:8], [20, '40.00%'])
        sheet = load_workbook(io.BytesIO(reports.xlsx(report))).active
        self.assertEqual(sheet.cell(1,4).value, '订单状态')
        self.assertEqual(sheet.cell(3,8).value, '40.00%')
        self.con.execute('UPDATE order_lines SET price_cents=0')
        self.assertIsNone(reports.build(self.con, {'kind':'orders'})['rows'][0][7])

    def test_filtered_summaries_and_export(self):
        self.assertEqual(reports.build(self.con, {'kind':'orders'})['summary'][4:7], [0,0,0])
        self.con.execute('UPDATE order_lines SET price_cents=1000')
        dm.create_purchase(self.con, dict(platform='淘宝', platform_order='32345678901', shop='店铺', amount='10', lines=[dict(order_line_id=self.new, quantity=3, cost='10')]))
        report = reports.build(self.con, {'kind':'orders'})
        self.assertEqual(report['summary'][4:8], [50,30,20,'40.00%'])
        self.assertEqual(reports.build(self.con, {'kind':'orders','customer':'不存在'})['summary'][4], 0)
        invoices = reports.build(self.con, {'kind':'invoices'})
        self.assertEqual(invoices['summary'][6:9], [30,0,30])
        sheet = load_workbook(io.BytesIO(reports.xlsx(invoices))).active
        self.assertEqual(sheet.cell(2,1).value, '汇总')
        self.assertEqual(sheet.cell(2,7).value, 30)
        self.assertEqual(reports.build(self.con, {'kind':'invoices','customer':'不存在'})['summary'][6:9], [0,0,0])
