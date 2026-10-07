import io
import unittest
from openpyxl import load_workbook
from pypdf import PdfReader
from server.exports import xlsx_quote, pdf_quote


class QuoteExportTests(unittest.TestCase):
    def test_formats_preserve_text_prices_and_all_lines(self):
        d = {'company': {'name': '测试公司', 'address': '上海', 'phone': '123'},
             'customer': '测试客户', 'po': '0000123', 'version': 2, 'contact': '联系人',
             'phone': '456', 'note': '含税报价', 'total_cents': 12345 * 54,
             'lines': [{'name': '=料品', 'spec': '规格型号', 'brand': '品牌',
                        'description': '描述不应导出', 'project_code': 'P301581', 'subproject_code': 'SUB-002', 'remark': '逐项备注', 'quantity': 3, 'unit': '个',
                        'price_cents': 4115, 'amount_cents': 12345}] * 54}
        ws = load_workbook(io.BytesIO(xlsx_quote(d))).active
        self.assertEqual(ws['B7'].data_type, 's')
        self.assertEqual(ws['I6'].value, '项目号')
        self.assertEqual(ws['I7'].value, 'P301581\nSUB-002')
        self.assertTrue(ws['I7'].alignment.wrap_text)
        self.assertEqual(ws['J6'].value, '备注')
        self.assertEqual(ws['J7'].value, '逐项备注')
        self.assertNotIn('描述不应导出', str(list(ws.values)))
        self.assertEqual(ws['G7'].value, 41.15)
        self.assertEqual(sum(1 for row in ws.iter_rows() if row[1].value == '=料品'), 54)
        self.assertTrue(any(row[7].value == 6666.3 for row in ws.iter_rows()))
        self.assertIn('&P', ws.oddFooter.center.text)
        self.assertGreater(len(ws.row_breaks.brk), 0)
        pdf = PdfReader(io.BytesIO(pdf_quote(d)))
        text = ''.join(page.extract_text() for page in pdf.pages)
        self.assertGreater(len(pdf.pages), 1)
        self.assertEqual(text.count('=料品'), 54)
        self.assertIn('6666.30', text)
        self.assertIn('0000123', text)
        self.assertIn('项目号', text)
        self.assertEqual(text.count('P301581'), 54)
        self.assertEqual(text.count('SUB-002'), 54)
        self.assertEqual(text.count('逐项备注'), 54)
        self.assertNotIn('描述不应导出', text)


if __name__ == '__main__':
    unittest.main()
