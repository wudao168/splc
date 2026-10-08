import base64
import http.client
import io
import json
import os
import sqlite3
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
import zipfile
from pathlib import Path
from http.server import ThreadingHTTPServer
from unittest.mock import patch
from server import app
from server.db import connect


class QuietHandler(app.Handler):
    def log_message(self, *_):
        pass


class HttpTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.old_data, self.old_connect = app.DATA, app.connect
        app.DATA = Path(self.temp.name)
        app.connect = lambda: connect(self.temp.name)
        self.server = ThreadingHTTPServer(('127.0.0.1',0), QuietHandler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.url = f'http://127.0.0.1:{self.server.server_port}'
        request = urllib.request.Request(self.url + '/api/auth/login',
            data=json.dumps({'username':'admin','password':'11111111'}).encode(),
            headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(request) as response:
            self.cookie = response.headers['Set-Cookie'].split(';', 1)[0]

    def tearDown(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()
        app.DATA, app.connect = self.old_data, self.old_connect
        self.temp.cleanup()

    def post(self, path, payload, origin=None):
        headers = {'Content-Type':'application/json', 'Cookie':self.cookie}
        if origin:
            headers['Origin'] = origin
        req = urllib.request.Request(self.url + path, data=json.dumps(payload).encode(), headers=headers)
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.load(r)
        except urllib.error.HTTPError as ex:
            return ex.code, json.load(ex)

    def get(self, path):
        return urllib.request.urlopen(urllib.request.Request(self.url + path, headers={'Cookie':self.cookie}))

    def request(self, method, path, body=None, headers=None):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=15)
        try:
            connection.request(method, path, body=body, headers=headers or {})
            response = connection.getresponse()
            payload = response.read()
            return response.status, json.loads(payload) if payload else {}
        finally:
            connection.close()

    def test_user_column_settings_persist_and_validate(self):
        self.assertEqual(self.post('/api/users/me/columns', {'table': 'purchases', 'hidden': ['logistics', 'platform_order']})[0], 200)
        with self.get('/api/auth/status') as response:
            user = json.load(response)['user']
        self.assertEqual(user['column_settings']['purchases'], ['logistics', 'platform_order'])
        self.assertEqual(self.post('/api/users/me/columns', {'table': 'unknown', 'hidden': []}), (400, {'error': '列设置对象无效'}))
        self.assertEqual(self.post('/api/users/me/columns', {'table': 'orders', 'hidden': ['bad key!']})[0], 400)
        self.assertEqual(self.post('/api/users/me/columns', {'table': 'orders', 'hidden': ['note']})[0], 200)
        with self.get('/api/auth/status') as response:
            user = json.load(response)['user']
        self.assertEqual(user['column_settings'], {'purchases': ['logistics', 'platform_order'], 'orders': ['note']})
        self.assertEqual(self.post('/api/users/me/columns', {'table': 'orders', 'hidden': []})[0], 200)
        with self.get('/api/auth/status') as response:
            user = json.load(response)['user']
        self.assertEqual(user['column_settings'], {'purchases': ['logistics', 'platform_order']})

    def test_products_purchases_and_inventory_http_routes(self):
        """产品库建档、仅关联料品的采购、入库、型号对应与出库都要能通过 HTTP 接口完成。"""
        self.assertEqual(self.post('/api/items', {'name': 'HTTP 料品', 'spec': 'A-1', 'unit': '个'})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        item = next(x for x in state['items'] if x['name'] == 'HTTP 料品')
        self.assertEqual(item['code'], 'LP000001')

        customer = self.post('/api/customers', {'name': 'HTTP 客户'})[1]
        order = self.post('/api/orders', {'customer_id': customer['id'], 'customer': 'HTTP 客户', 'po': 'PO-HTTP',
            'address': '测试地址', 'lines': [{'name': 'HTTP 料品', 'spec': 'A-1', 'quantity': 5, 'unit': '个', 'price': 10}]})[1]
        self.assertEqual(self.post(f"/api/orders/{order['id']}/confirm", {})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        line = next(x for x in state['order_lines'] if x['order_id'] == order['id'])
        self.assertEqual(line['item_id'], item['id'], '订单明细应自动匹配同一料品编号')

        purchase = self.post('/api/purchases', {'platform': '京东', 'shop': '备货店铺', 'platform_order': 'HTTP-STOCK',
            'amount': 90, 'lines': [{'item_id': item['id'], 'quantity': 3, 'cost': 90, 'receive_mode': 'stock'}]})[1]
        with self.get('/api/state') as response:
            state = json.load(response)
        purchase_line = next(x for x in state['purchase_lines'] if x['purchase_id'] == purchase['id'])
        self.assertIsNone(purchase_line['order_line_id'])
        self.assertEqual(self.post(f"/api/purchase-lines/{purchase_line['id']}/receive-mode", {'receive_mode': 'direct'})[0], 200)
        self.assertEqual(self.post(f"/api/purchase-lines/{purchase_line['id']}/receive-mode", {'receive_mode': 'stock', 'warehouse_id': state['warehouses'][0]['id']})[0], 200)
        receipt = self.post('/api/receipts', {'warehouse_id': state['warehouses'][0]['id'],
            'lines': [{'purchase_line_id': purchase_line['id'], 'quantity': 3}]})[1]
        self.assertTrue(receipt['number'].startswith('RKD-'))
        self.assertEqual(self.post(f"/api/purchase-lines/{purchase_line['id']}/receive-mode", {'receive_mode': 'direct'})[0], 400)
        self.assertEqual(self.post(f"/api/items/{item['id']}/aliases", {'source': '客户型号', 'alias': 'ABC-100'})[0], 200)
        outbound = self.post('/api/outbounds', {'warehouse_id': state['warehouses'][0]['id'], 'order_id': order['id'],
            'company': '测试公司', 'lines': [{'order_line_id': line['id'], 'quantity': 2}]})[1]
        self.assertTrue(outbound['delivery_id'])
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(next(x for x in state['items'] if x['id'] == item['id'])['on_hand'], 1)
        self.assertEqual(len(state['stock_reservations']), 1)
        self.assertEqual(state['order_lines'][0]['stock_cost_cents'], 6000)
        self.assertEqual(self.post('/api/items/backfill', {})[0], 200)
        self.assertEqual(self.post('/api/items', {'name': ''})[0], 400)

    def test_lan_hosts_are_allowed_only_when_known(self):
        for method, path, body in (('GET', '/api/health', None), ('POST', '/api/auth/logout', b'{}')):
            headers = {'Host': 'evil.example'} if method == 'GET' else {'Host': 'evil.example', 'Content-Type': 'application/json'}
            self.assertEqual(self.request(method, path, body, headers), (403, {'error': '仅允许本机访问'}))

        lan = frozenset({'127.0.0.1', 'localhost', 'caidan.lan', '192.168.10.63'})
        login = json.dumps({'username': 'admin', 'password': '11111111'}).encode()
        with patch.object(app, 'ALLOWED_HOSTS', lan):
            self.assertEqual(self.request('GET', '/api/health', headers={'Host': 'caidan.lan:8765'})[0], 200)
            self.assertEqual(self.request('GET', '/api/auth/status', headers={'Host': '192.168.10.63:8765'})[0], 200)
            self.assertEqual(self.request('POST', '/api/auth/login', login, {
                'Host': 'caidan.lan:8765', 'Content-Type': 'application/json', 'Origin': 'http://caidan.lan:8765'})[0], 200)
            self.assertEqual(self.request('POST', '/api/auth/login', login, {
                'Host': 'caidan.lan:8765', 'Content-Type': 'application/json', 'Origin': 'http://evil.example'}), (403, {'error': '来源不受信任'}))

    def test_public_ip_hosts_work_without_configuration(self):
        for host in ('8.8.8.8:8765', '1.1.1.1', '[2001:4860:4860::8888]:8765'):
            self.assertEqual(self.request('GET', '/api/health', headers={'Host': host})[0], 200)
        for host in ('10.99.99.99:8765', '203.0.113.9:8765', 'evil.example'):
            self.assertEqual(self.request('GET', '/api/health', headers={'Host': host}), (403, {'error': '仅允许本机访问'}))
        login = json.dumps({'username': 'admin', 'password': '11111111'}).encode()
        self.assertEqual(self.request('POST', '/api/auth/login', login, {
            'Host': '8.8.8.8:8765', 'Content-Type': 'application/json', 'Origin': 'http://8.8.8.8:8765'})[0], 200)

    def test_login_company_and_actor_audit(self):
        with self.assertRaises(urllib.error.HTTPError) as denied:
            urllib.request.urlopen(self.url + '/api/state')
        self.assertEqual(denied.exception.code, 401)

        self.assertEqual(self.post('/api/auth/login', {'username':'admin','password':'wrong'})[0], 400)
        self.assertEqual(self.post('/api/users/me/theme', {'theme':'o'}), (200, {'theme':'o'}))
        with self.get('/api/auth/status') as response:
            self.assertEqual(json.load(response)['user']['theme'], 'o')
        self.assertEqual(self.post('/api/users/me/theme', {'theme':'invalid'})[0], 400)
        self.assertEqual(self.post('/api/users', {'username':'second','display_name':'第二用户','password':'testing456'})[0], 200)
        self.assertEqual(self.post('/api/company', {'name':'测试公司','name_en':'Test Co','address':'测试地址','phone':'12345'})[0], 200)
        self.assertEqual(self.post('/api/activity', {'route':'settings'})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(state['company']['name'], '测试公司')
        self.assertEqual([user['username'] for user in state['users']], ['admin', 'second'])
        self.assertEqual(state['audit'][0]['user_id'], state['users'][0]['id'])
        self.assertEqual(self.post('/api/auth/logout', {})[0], 200)
        with self.assertRaises(urllib.error.HTTPError) as denied:
            self.get('/api/state')
        self.assertEqual(denied.exception.code, 401)

    def test_customer_invoice_info_requires_login_and_preserves_other_fields(self):
        code, customer = self.post('/api/customers', {'name': '开票客户', 'note': '保留备注'})
        self.assertEqual(code, 200)
        route = f"/api/customers/{customer['id']}/invoice-info"
        self.assertEqual(self.post(route, {'title': '开票客户抬头', 'bank_account': '000123'})[0], 200)
        self.assertEqual(self.post(route, {'tax_number': '未填写抬头'})[0], 400)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(state['customers'][0]['note'], '保留备注')
        self.assertEqual(state['customers'][0]['invoice_info']['bank_account'], '000123')
        self.assertEqual(state['audit'][0]['action'], '保存客户开票信息')
        self.assertIsNotNone(state['audit'][0]['user_id'])
        self.post('/api/auth/logout', {})
        self.assertEqual(self.post(route, {'title': '未登录修改'})[0], 401)

    def test_customer_invoice_exports_current_draft_without_saving(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        _, customer = self.post('/api/customers', {'name': '客户资料名称'})
        route = f"/api/customers/{customer['id']}/invoice-info"
        self.assertEqual(self.post(route, {'title': '已保存抬头', 'bank_account': '00123'})[0], 200)
        self.post('/api/company', {'name': '本公司资料'})
        with self.get('/api/state') as response:
            before = json.load(response)
        draft = {'title': '客户<开票>&抬头', 'tax_number': '001234567890123456', 'address': '客户注册地址 100 号',
                 'phone': '021-12345678', 'bank_name': '=1+1', 'bank_account': '00012345678901234567890', 'email': 'customer@example.com'}
        for format in ('pdf', 'xlsx'):
            request = urllib.request.Request(self.url + route + '/' + format, data=json.dumps(draft).encode(),
                                             headers={'Content-Type': 'application/json', 'Cookie': self.cookie})
            with urllib.request.urlopen(request) as response:
                blob = response.read()
                self.assertIn('attachment;', response.headers['Content-Disposition'])
            if format == 'xlsx':
                ws = load_workbook(io.BytesIO(blob)).active
                self.assertEqual(ws['A1'].value, '客户开票信息')
                for index, value in enumerate(draft.values(), 3):
                    self.assertEqual(ws.cell(index, 2).value, value)
                    self.assertEqual(ws.cell(index, 2).data_type, 's')
                    self.assertEqual(ws.cell(index, 2).number_format, '@')
            else:
                pdf = PdfReader(io.BytesIO(blob))
                self.assertEqual(len(pdf.pages), 1)
                text = pdf.pages[0].extract_text()
                self.assertIn('客户开票信息', text)
                for value in draft.values():
                    self.assertIn(value, text)
        with self.get('/api/state') as response:
            after = json.load(response)
        self.assertEqual(after['customers'], before['customers'])
        self.assertEqual(after['company'], before['company'])
        self.assertEqual(after['audit'], before['audit'])
        self.assertEqual(self.post(route + '/pdf', {})[0], 400)
        self.assertEqual(self.post('/api/customers/99999/invoice-info/pdf', draft)[0], 400)
        self.post('/api/auth/logout', {})
        self.assertEqual(self.post(route + '/xlsx', draft)[0], 401)

    def test_remember_login_persists_cookie_until_logout(self):
        request = urllib.request.Request(self.url + '/api/auth/login',
            data=json.dumps({'username':'admin','password':'11111111','remember':True}).encode(),
            headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(request) as response:
            cookie_header = response.headers['Set-Cookie']
        self.assertIn('Max-Age=31536000', cookie_header)
        self.assertIn('HttpOnly', cookie_header)
        self.assertNotIn('11111111', cookie_header)
        cookie = cookie_header.split(';', 1)[0]
        status = urllib.request.Request(self.url + '/api/auth/status', headers={'Cookie':cookie})
        with urllib.request.urlopen(status) as response:
            self.assertEqual(json.load(response)['user']['username'], 'admin')
        logout = urllib.request.Request(self.url + '/api/auth/logout', data=b'{}',
            headers={'Content-Type':'application/json', 'Cookie':cookie})
        with urllib.request.urlopen(logout) as response:
            self.assertIn('Max-Age=0', response.headers['Set-Cookie'])
        with urllib.request.urlopen(status) as response:
            self.assertIsNone(json.load(response)['user'])

    def test_billing_profile_exports_and_partial_company_update(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        profile = {'name': '测试智能科技有限公司', 'address': '测试地址1088弄7号', 'phone': '18551211185',
                   'bank_name': '测试银行', 'bank_account': '0011291909300214567',
                   'tax_number': '91310120MADL4XAW73', 'email': 'test@example.com'}
        self.assertEqual(self.post('/api/company', profile)[0], 200)
        self.assertEqual(self.post('/api/company', {'name': '更新公司名称'})[0], 200)
        with self.get('/api/state') as response:
            saved = json.load(response)['company']
        self.assertEqual(saved['bank_account'], profile['bank_account'])
        self.assertEqual(saved['email'], profile['email'])
        self.assertEqual(saved['name'], '更新公司名称')

        draft = {**profile, 'bank_name': '=1+1'}
        for format in ('pdf', 'xlsx'):
            request = urllib.request.Request(self.url + '/api/company/billing/' + format,
                data=json.dumps(draft).encode(), headers={'Content-Type': 'application/json', 'Cookie': self.cookie})
            with urllib.request.urlopen(request) as response:
                blob = response.read()
                self.assertIn('attachment;', response.headers['Content-Disposition'])
            if format == 'xlsx':
                ws = load_workbook(io.BytesIO(blob)).active
                self.assertEqual(ws['B3'].value, profile['name'])
                self.assertEqual(ws['B4'].value, '=1+1')
                self.assertEqual(ws['B4'].data_type, 's')
                self.assertEqual(ws['B5'].value, profile['bank_account'])
                self.assertEqual(ws['B5'].number_format, '@')
            else:
                pdf = PdfReader(io.BytesIO(blob))
                self.assertEqual(len(pdf.pages), 1)
                text = pdf.pages[0].extract_text()
                self.assertIn(profile['name'], text)
                self.assertIn(profile['bank_account'], text)
        with self.get('/api/state') as response:
            self.assertEqual(json.load(response)['company'], saved)
        self.assertEqual(self.post('/api/company/billing/pdf', {'name': ''})[0], 400)
        self.assertEqual(self.post('/api/company', {'email': 'x' * 501})[0], 400)
        request = urllib.request.Request(self.url + '/api/company/billing/pdf', data=b'{}', headers={'Content-Type':'application/json'})
        with self.assertRaises(urllib.error.HTTPError) as denied:
            urllib.request.urlopen(request)
        self.assertEqual(denied.exception.code, 401)

    def test_order_salesperson_selection(self):
        self.assertEqual(self.post('/api/company', {'name': '测试公司', 'company_code': 'SL'})[0], 200)
        created = self.post('/api/users', {'username': 'sales01', 'display_name': '业务员甲', 'password': '11111111'})
        self.assertEqual(created[0], 200)
        lines = [{'name': '料品', 'quantity': 1, 'unit': '个', 'price': 10}]
        self.assertEqual(self.post('/api/orders', {'customer': '业务员客户', 'po': 'PO-SALES', 'address': '地址',
            'salesperson_id': created[1]['id'], 'lines': lines})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        order = [o for o in state['orders'] if o['po'] == 'PO-SALES'][0]
        self.assertEqual(order['salesperson_id'], created[1]['id'])
        self.assertEqual([u['display_name'] for u in state['users'] if u['id'] == created[1]['id']], ['业务员甲'])
        # 留空表示未指定，非法编号被拦截
        _, plain = self.post('/api/orders', {'customer': '业务员客户二', 'po': 'PO-SALES-2', 'address': '地址',
            'salesperson_id': '', 'lines': lines})
        with self.get('/api/state') as response:
            self.assertIsNone([o for o in json.load(response)['orders'] if o['id'] == plain['id']][0]['salesperson_id'])
        self.assertEqual(self.post('/api/orders', {'customer': '业务员客户三', 'po': 'PO-SALES-3', 'address': '地址',
            'salesperson_id': 999, 'lines': lines}), (400, {'error': '业务员不存在，请重新选择'}))
        self.assertEqual(self.post('/api/orders', {'customer': '业务员客户四', 'po': 'PO-SALES-4', 'address': '地址',
            'salesperson_id': '甲', 'lines': lines}), (400, {'error': '业务员编号无效'}))
        # 修改订单时未提交业务员则保留原值
        self.assertEqual(self.post(f"/api/orders/{order['id']}/edit", {'customer': order['customer'], 'po': order['po'], 'contact': '',
            'phone': '', 'address': '地址', 'due_date': '', 'note': '', 'lines': lines})[0], 200)
        with self.get('/api/state') as response:
            self.assertEqual([o for o in json.load(response)['orders'] if o['id'] == order['id']][0]['salesperson_id'], created[1]['id'])

    def test_auto_internal_po_number_rule_and_sequence(self):
        from datetime import datetime
        from server.domain import TZ
        day = datetime.now(TZ).strftime('%Y%m%d')
        self.assertEqual(self.post('/api/company', {'name': '测试公司', 'company_code': 'sl'})[0], 200)
        with self.get('/api/state') as response:
            self.assertEqual(json.load(response)['company']['company_code'], 'SL')
        self.assertEqual(self.post('/api/orders/next-po', {})[1],
                         {'po': f'POSL{day}001', 'company_code': 'SL', 'date': day})
        self.assertEqual(self.post('/api/orders/next-po', {})[1]['po'], f'POSL{day}001')  # 预览不占用流水号
        payload = {'customer': '自动编号客户', 'address': '测试地址', 'auto_po': True, 'po': f'POSL{day}001',
                   'lines': [{'name': '料品', 'quantity': 1, 'unit': '个', 'price': 10}]}
        self.assertEqual(self.post('/api/orders', payload)[0], 200)
        self.assertEqual(self.post('/api/orders/next-po', {})[1]['po'], f'POSL{day}002')
        untitled = {'customer': '自动编号客户二', 'address': '测试地址', 'auto_po': True, 'po': '',
                    'lines': [{'name': '料品', 'quantity': 1, 'unit': '个', 'price': 10}]}
        self.assertEqual(self.post('/api/orders', untitled)[0], 200)
        with self.get('/api/state') as response:
            pos = [order['po'] for order in json.load(response)['orders']]
        self.assertEqual(sorted(pos), sorted({f'POSL{day}001', f'POSL{day}002'}))
        self.assertEqual(self.post('/api/company', {'company_code': ''})[0], 200)
        self.assertEqual(self.post('/api/orders/next-po', {}),
                         (400, {'error': '请先在“设置 · 公司信息”中填写公司代码（一般为 2 个字母）'}))
        self.assertEqual(self.post('/api/company', {'company_code': 'A1B'})[0], 400)

    def test_invoice_reminder_and_platform_apply_queue(self):
        """催票登记与「平台内申请开票」任务队列。"""
        _, order = self.post('/api/orders', {'customer': '催票客户', 'po': 'PO-REMIND', 'address': '地址',
            'lines': [{'name': '料品', 'quantity': 2, 'price': 100}]})
        self.assertEqual(self.post(f"/api/orders/{order['id']}/confirm", {})[0], 200)
        with self.get('/api/state') as response:
            line_id = next(l['id'] for l in json.load(response)['order_lines'] if l['order_id'] == order['id'])
        _, purchase = self.post('/api/purchases', {'platform': '淘宝', 'shop': '催票店铺', 'platform_order': 'TAO-REMIND-1',
            'amount': 200, 'lines': [{'order_line_id': line_id, 'quantity': 2}]})
        pid = purchase['id']
        code, reminded = self.post(f'/api/purchases/{pid}/invoice-remind', {'channel': '聊天催票', 'note': '已联系商家'})
        self.assertEqual((code, reminded['invoice_remind_count']), (200, 1))
        self.assertTrue(reminded['next_followup'])
        code, queued = self.post(f'/api/purchases/{pid}/apply-invoice', {})
        self.assertEqual((code, queued['status']), (200, 'pending'))
        self.assertEqual(self.post(f'/api/purchases/{pid}/apply-invoice', {})[0], 400)
        with self.get('/api/invoice-apply/pending') as response:
            requests = json.load(response)['requests']
        self.assertEqual([r['platform_order'] for r in requests], ['TAO-REMIND-1'])
        with patch.dict(os.environ, {'CAIDAN_INTERNAL_TOKEN': 'unit-test-token'}):
            request = urllib.request.Request(self.url + f"/api/invoice-apply/{queued['id']}/result",
                data=json.dumps({'status': 'done', 'message': '已点击申请开票'}).encode(),
                headers={'Content-Type': 'application/json', 'X-Caidan-Internal': 'unit-test-token'})
            with urllib.request.urlopen(request) as response:
                self.assertEqual(json.load(response)['status'], 'done')
            self.assertEqual(self.post(f"/api/invoice-apply/{queued['id']}/result", {'status': 'unknown'})[0], 400)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(state['purchases'][0]['invoice_remind_count'], 2)
        self.assertEqual(state['invoice_apply_requests'][0]['status'], 'done')

    def test_company_stamps_save_replace_remove_and_validate_image(self):
        from PIL import Image
        ids = []
        for color in ('red', 'blue'):
            out = io.BytesIO()
            Image.new('RGBA', (24, 24), color).save(out, format='PNG')
            status, result = self.post('/api/attachments', {'name': color + '.png', 'content': base64.b64encode(out.getvalue()).decode()})
            self.assertEqual(status, 200)
            ids.append(result['id'])
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'contract', 'attachment_id': ids[0]})[0], 200)
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'delivery', 'attachment_id': ids[1]})[0], 200)
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'contract', 'attachment_id': ids[1]})[0], 200)
        with self.get('/api/state') as response:
            self.assertEqual(json.load(response)['company_stamps'], {'contract': ids[1], 'delivery': ids[1]})
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'contract', 'attachment_id': ''})[0], 200)
        with self.get('/api/state') as response:
            self.assertEqual(json.load(response)['company_stamps'], {'delivery': ids[1]})
        with self.get('/api/files/' + ids[1]) as response:
            self.assertTrue(response.read().startswith(b'\x89PNG'))
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'invalid', 'attachment_id': ids[0]})[0], 400)
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'delivery', 'attachment_id': 'missing'})[0], 400)
        _, fake = self.post('/api/attachments', {'name': 'fake.png', 'content': base64.b64encode(b'not an image').decode()})
        self.assertEqual(self.post('/api/company/stamps', {'kind': 'delivery', 'attachment_id': fake['id']})[0], 400)

    def test_ordinary_user_can_cancel_order_and_record_purchase_handling(self):
        self.assertEqual(self.post('/api/users', {'username':'staff','display_name':'普通用户','password':'testing456'})[0], 200)
        _, order = self.post('/api/orders', {'customer':'测试客户','po':'PO-CANCEL','address':'测试地址',
            'lines':[{'name':'螺栓','quantity':1}]})
        self.assertEqual(self.post(f"/api/orders/{order['id']}/confirm", {})[0], 200)
        with self.get('/api/state') as response:
            line_id = next(l['id'] for l in json.load(response)['order_lines'] if l['order_id'] == order['id'])
        _, purchase = self.post('/api/purchases', {'platform':'京东','shop':'测试店铺',
            'platform_order':'JD-CANCEL','amount':10,'lines':[{'order_line_id':line_id,'quantity':1}]})
        with self.get('/api/state') as response:
            purchase_line_id = next(l['id'] for l in json.load(response)['purchase_lines'] if l['purchase_id'] == purchase['id'])
        request = urllib.request.Request(self.url + '/api/auth/login',
            data=json.dumps({'username':'staff','password':'testing456'}).encode(),
            headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(request) as response:
            self.cookie = response.headers['Set-Cookie'].split(';', 1)[0]
        self.assertEqual(self.post('/api/orders/delete', {'items':[{'type':'order','id':order['id']}]})[0], 400)
        self.assertEqual(self.post(f"/api/orders/{order['id']}/cancel", {'reason':'客户取消'})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        case = state['purchase_cases'][0]
        self.assertEqual(self.post(f"/api/purchase-cases/{case['id']}/update", {'kind':'cancel','amount':10,'note':'已联系卖家退款'})[0], 200)
        self.assertEqual(self.post(f"/api/purchase-cases/{case['id']}/finance", {})[0], 400)
        self.assertEqual(self.post(f"/api/purchase-cases/{case['id']}/complete", {})[0], 400)
        self.assertEqual(state['orders'][0]['status'], 'cancelled')
        self.assertEqual(state['purchase_cases'][0]['quantity'], 1)

    def test_internal_invoice_attachment_upload(self):
        content = base64.b64encode(b'%PDF-1.4 test invoice').decode()
        request = urllib.request.Request(self.url + '/api/attachments',
            data=json.dumps({'name':'invoice.pdf','content':content}).encode(),
            headers={'Content-Type':'application/json', 'X-Caidan-Internal':'test-token'})
        with patch.dict(os.environ, {'CAIDAN_INTERNAL_TOKEN':'test-token'}):
            with urllib.request.urlopen(request) as response:
                self.assertEqual(response.status, 200)
                attachment = json.load(response)
        with self.get('/api/state') as response:
            self.assertEqual(json.load(response)['attachments'][0]['id'], attachment['id'])

    def test_rollback_on_partial_invalid_order(self):
        code, _ = self.post('/api/orders', {'customer':'测试','po':'X','address':'测试地址','lines':[
            {'name':'合法料品','quantity':1}, {'name':'非法数量','quantity':-2}]})
        self.assertEqual(code,400)
        with self.get('/api/state') as r:
            state = json.load(r)
        self.assertEqual(state['orders'],[])
        self.assertEqual(state['order_lines'],[])

    def test_purchase_attachment_endpoints(self):
        order = self.post('/api/orders', {'customer':'测试','po':'FILES','address':'地址',
            'lines':[{'name':'螺栓','quantity':2}]})[1]
        self.assertEqual(self.post(f"/api/orders/{order['id']}/confirm", {})[0], 200)
        files = [self.post('/api/attachments', {'name':name,'content':base64.b64encode(raw).decode()})[1]
                 for name,raw in [('one.pdf', b'%PDF-one'), ('two.png', b'PNG-two')]]
        with self.get('/api/state') as response:
            line_id = json.load(response)['order_lines'][0]['id']
        code, purchase = self.post('/api/purchases', {'platform':'京东','shop':'店铺','platform_order':'FILES',
            'amount':10,'attachment_ids':[files[0]['id']],
            'lines':[{'order_line_id':line_id,'quantity':1}]})
        self.assertEqual(code, 200)
        pid = purchase['id']
        self.assertEqual(self.post(f'/api/purchases/{pid}/attachments', {'attachment_ids':[files[1]['id']]})[0], 200)
        self.assertEqual(self.post(f'/api/purchases/{pid}/attachments/remove', {'attachment_id':files[0]['id']})[0], 200)
        with self.get('/api/state') as response:
            saved = json.load(response)['purchases'][0]
        self.assertEqual([file['name'] for file in saved['attachments']], ['two.png'])
        self.assertEqual(saved['source_id'], files[1]['id'])

    def test_untrusted_origin_and_backup_restore(self):
        payload = {'customer':'测试','po':'X','address':'测试地址','lines':[{'name':'合法料品','quantity':1}]}
        self.assertEqual(self.post('/api/orders',payload,'https://untrusted.example')[0],403)
        self.assertEqual(self.post('/api/orders',payload)[0],200)
        self.assertEqual(self.post('/api/orders',payload)[0],400)
        content = base64.b64encode(b'%PDF-1.4 TEST').decode()
        code, file = self.post('/api/attachments',{'name':'invoice.pdf','content':content})
        self.assertEqual(code,200)
        with self.get('/api/backup') as r:
            archive = zipfile.ZipFile(io.BytesIO(r.read()))
        restored = Path(self.temp.name)/'restored.sqlite3'
        restored.write_bytes(archive.read('caidan.sqlite3'))
        con = sqlite3.connect(restored)
        try:
            self.assertEqual(con.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertEqual(con.execute('SELECT count(*) FROM orders').fetchone()[0],1)
        finally:
            con.close()
        self.assertIn('attachments/'+file['id'],archive.namelist())


    def test_repeated_preview_is_not_saved_order(self):
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active
        sheet.append(['料品名称', '数量']); sheet.append(['螺栓', 2])
        output = io.BytesIO(); book.save(output)
        payload = {'name': '询价.xlsx', 'content': base64.b64encode(output.getvalue()).decode()}
        code, first = self.post('/api/imports', payload)
        self.assertEqual(code, 200)
        self.assertFalse(first['duplicate'])
        code, second = self.post('/api/imports', payload)
        self.assertEqual(code, 200)
        self.assertFalse(second['duplicate'])
        self.assertEqual(second['existing_orders'], [])
        code, _ = self.post('/api/orders', {'customer':'测试', 'po':'SAVED-001', 'address':'地址',
            'source_id':first['source_id'], 'lines':[{'name':'螺栓','quantity':2}]})
        self.assertEqual(code, 200)
        code, third = self.post('/api/imports', payload)
        self.assertTrue(third['duplicate'])
        self.assertEqual(third['existing_orders'][0]['po'], 'SAVED-001')


    def test_import_file_size_boundary_and_attachment_limit(self):
        # Exercise the actual HTTP/base64/file-storage boundary; parsing is covered separately.
        parsed = {'tables': [], 'text': '', 'metadata': {}, 'warnings': []}
        with patch('server.app.parse_file', return_value=parsed) as parser:
            payload = {'name': 'large.pdf', 'content': base64.b64encode(b'x' * (100 * 1024 * 1024)).decode()}
            code, result = self.post('/api/imports', payload)
            self.assertEqual(code, 200, result)
            parser.assert_called_once()
            self.assertEqual((app.DATA / 'attachments' / result['source_id']).stat().st_size, 100 * 1024 * 1024)
            payload['content'] = base64.b64encode(b'x' * (100 * 1024 * 1024 + 1)).decode()
            code, error = self.post('/api/imports', payload)
            self.assertEqual(code, 400)
            self.assertIn('100 MB', error['error'])
            self.assertEqual(parser.call_count, 1)
            del payload
        payload = {'name': 'attachment.pdf', 'content': base64.b64encode(b'x' * (21 * 1024 * 1024)).decode()}
        code, error = self.post('/api/attachments', payload)
        self.assertEqual(code, 400)
        self.assertIn('20 MB', error['error'])
        with self.get('/api/state') as response:
            self.assertEqual(len(json.load(response)['attachments']), 1)

    def test_incomplete_intake_draft_can_resume_and_complete(self):
        payload = {'form': {'customer_id': '', 'po': '', 'address': ''},
                   'lines': [{'name': '', 'quantity': '', 'unit': '个'}], 'selected': 0}
        code, saved = self.post('/api/intake-drafts', {'payload': payload})
        self.assertEqual(code, 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(state['intake_drafts'][0]['payload'], payload)
        self.assertEqual(state['orders'], [])
        updated = {**payload, 'form': {**payload['form'], 'po': 'PO-READY'}}
        code, same = self.post(f"/api/intake-drafts/{saved['id']}", {'payload': updated})
        self.assertEqual((code, same['id']), (200, saved['id']))
        order = {'customer':'测试', 'po':'PO-READY', 'address':'测试地址',
                 'intake_draft_id':saved['id'], 'lines':[{'name':'螺栓','quantity':1,'unit':'个'}]}
        self.assertEqual(self.post('/api/orders', order)[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(state['intake_drafts'], [])
        self.assertEqual(state['orders'][0]['po'], 'PO-READY')

    def test_delete_intake_draft_only_hides_selected_draft(self):
        payload = {'form': {'po': 'PO-DELETE'}, 'lines': [{'name': '螺栓', 'quantity': 1}]}
        _, deleted = self.post('/api/intake-drafts', {'payload': payload})
        _, kept = self.post('/api/intake-drafts', {'payload': {**payload, 'form': {'po': 'PO-KEEP'}}})
        self.assertEqual(self.post(f"/api/intake-drafts/{deleted['id']}/delete", {})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual([x['id'] for x in state['intake_drafts']], [kept['id']])
        self.assertEqual(state['orders'], [])
        self.assertEqual(self.post(f"/api/intake-drafts/{deleted['id']}", {'payload': payload})[0], 400)
        self.assertEqual(self.post(f"/api/intake-drafts/{deleted['id']}/delete", {})[0], 400)

    def test_batch_delete_orders_is_atomic_and_keeps_purchases(self):
        draft_payload = {'form': {'po': 'DRAFT-1'}, 'lines': [{'name': '螺栓', 'quantity': 1}]}
        _, draft = self.post('/api/intake-drafts', {'payload': draft_payload})
        order_payload = lambda po: {'customer': '测试客户', 'po': po, 'address': '测试地址',
                                    'lines': [{'name': '螺栓', 'quantity': 1}]}
        _, removable = self.post('/api/orders', order_payload('PO-REMOVE'))
        self.assertEqual(self.post(f"/api/orders/{removable['id']}/confirm", {})[0], 200)
        _, linked = self.post('/api/orders', order_payload('PO-LINKED'))
        self.assertEqual(self.post(f"/api/orders/{linked['id']}/confirm", {})[0], 200)
        with self.get('/api/state') as response:
            state = json.load(response)
        line_id = next(line['id'] for line in state['order_lines'] if line['order_id'] == linked['id'])
        self.assertEqual(self.post('/api/purchases', {'platform': '京东', 'shop': '测试店铺',
            'platform_order': 'JD-DELETE-GUARD', 'amount': '10',
            'lines': [{'order_line_id': line_id, 'quantity': 1, 'purchase_quantity': 1, 'cost': '10'}]})[0], 200)
        items = [{'type': 'intake', 'id': draft['id']}, {'type': 'order', 'id': removable['id']},
                 {'type': 'order', 'id': linked['id']}]
        code, error = self.post('/api/orders/delete', {'items': items})
        self.assertEqual(code, 400)
        self.assertIn('已关联采购', error['error'])
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(len(state['orders']), 2)
        self.assertEqual(len(state['intake_drafts']), 1)
        self.assertEqual(self.post('/api/orders/delete', {'items': items[:2]})[1]['deleted'], 2)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual([order['id'] for order in state['orders']], [linked['id']])
        self.assertEqual(state['intake_drafts'], [])
        self.assertEqual([quote['order_id'] for quote in state['quotes']], [linked['id']])
        self.assertEqual([line['order_id'] for line in state['order_lines']], [linked['id']])
        self.assertEqual(self.post('/api/orders', order_payload('PO-REMOVE'))[0], 400)
        self.assertEqual(self.post(f"/api/orders/{removable['id']}/restore", {})[0], 200)
        self.assertEqual(self.post('/api/orders/delete', {'items': [items[2]]})[0], 400)

    def test_admin_deletes_purchase_and_then_customer_order(self):
        _, order = self.post('/api/orders', {'customer':'测试客户', 'po':'PO-ADMIN-DELETE',
            'address':'测试地址', 'lines':[{'name':'螺栓', 'quantity':1}]})
        self.assertEqual(self.post(f"/api/orders/{order['id']}/confirm", {})[0], 200)
        with self.get('/api/state') as response:
            line_id = json.load(response)['order_lines'][0]['id']
        _, purchase = self.post('/api/purchases', {'platform':'淘宝', 'shop':'测试店铺',
            'platform_order':'TB-ADMIN-DELETE', 'amount':'10',
            'lines':[{'order_line_id':line_id, 'quantity':1, 'cost':'10'}]})
        with self.get('/api/state') as response:
            purchase_line_id = json.load(response)['purchase_lines'][0]['id']
        self.assertEqual(self.post('/api/orders/delete', {'items':[{'type':'order', 'id':order['id']}]})[0], 400)
        self.assertEqual(self.post('/api/purchases/delete', {'ids':[purchase['id']]}), (200, {'deleted':1}))
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual((state['purchases'], state['purchase_lines'], state['packages']), ([], [], []))
        self.assertEqual(state['order_lines'][0]['purchased'], 0)
        self.assertEqual(self.post('/api/orders/delete', {'items':[{'type':'order', 'id':order['id']}]}), (200, {'deleted':1}))

    def test_ordinary_user_deletes_unused_records_but_cannot_restore(self):
        self.assertEqual(self.post('/api/users', {'username':'worker', 'display_name':'采购员', 'password':'testing456'})[0], 200)
        _, order = self.post('/api/orders', {'customer':'测试客户','po':'PO-STAFF-DELETE','address':'测试地址','lines':[{'name':'螺栓','quantity':1}]})
        self.post(f"/api/orders/{order['id']}/confirm", {})
        with self.get('/api/state') as response:
            lid = json.load(response)['order_lines'][0]['id']
        _, purchase = self.post('/api/purchases', {'platform':'京东','shop':'店铺','platform_order':'STAFF-DELETE','amount':10,'lines':[{'order_line_id':lid,'quantity':1}]})
        request = urllib.request.Request(self.url + '/api/auth/login', data=json.dumps({'username':'worker','password':'testing456'}).encode(), headers={'Content-Type':'application/json'})
        admin_cookie = self.cookie
        with urllib.request.urlopen(request) as response:
            self.cookie = response.headers['Set-Cookie'].split(';', 1)[0]
        self.assertEqual(self.post('/api/orders/delete', {'items':[{'type':'order','id':order['id']}]})[0], 400)
        self.assertEqual(self.post('/api/purchases/delete', {'ids':[purchase['id']]}), (200, {'deleted':1}))
        self.assertEqual(self.post('/api/orders/delete', {'items':[{'type':'order','id':order['id']}]}), (200, {'deleted':1}))
        self.assertEqual(self.post(f"/api/orders/{order['id']}/restore", {})[0], 400)
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(len(state['trash']), 2)
        self.assertEqual(state['orders'], [])
        self.assertTrue(any(a['user_id'] != 1 and a['action'] == '移入回收站' for a in state['audit']))
        self.cookie = admin_cookie
        self.assertEqual(self.post(f"/api/orders/{order['id']}/restore", {})[0], 200)
        self.assertEqual(self.post(f"/api/purchases/{purchase['id']}/restore", {})[0], 200)

    def test_delete_preserves_delivery_history(self):
        _, order = self.post('/api/orders', {'customer':'测试客户', 'po':'PO-DELIVERY-GUARD',
            'address':'测试地址', 'lines':[{'name':'螺栓', 'quantity':1}]})
        self.assertEqual(self.post(f"/api/orders/{order['id']}/confirm", {})[0], 200)
        with self.get('/api/state') as response:
            line_id = json.load(response)['order_lines'][0]['id']
        _, purchase = self.post('/api/purchases', {'platform':'淘宝', 'shop':'测试店铺',
            'platform_order':'TB-DELIVERY-GUARD', 'amount':'10',
            'lines':[{'order_line_id':line_id, 'quantity':1, 'cost':'10'}]})
        with self.get('/api/state') as response:
            purchase_line_id = json.load(response)['purchase_lines'][0]['id']
        _, package = self.post('/api/packages', {'carrier':'测试快递', 'tracking':'DELIVERY-001',
            'lines':[{'purchase_line_id':purchase_line_id, 'quantity':1}]})
        self.assertEqual(self.post(f"/api/packages/{package['id']}/tracking", {'status':'运输中'})[0], 200)
        with self.get('/api/state') as response:
            package_line_id = json.load(response)['package_lines'][0]['id']
        self.assertEqual(self.post('/api/deliveries', {'company':'测试公司',
            'lines':[{'package_line_id':package_line_id, 'quantity':1}]})[0], 200)
        code, result = self.post('/api/purchases/delete', {'ids':[purchase['id']]})
        self.assertEqual(code, 400)
        self.assertTrue('包裹' in result['error'] or '送货' in result['error'])
        with self.get('/api/state') as response:
            state = json.load(response)
        self.assertEqual(len(state['purchases']), 1)
        self.assertEqual(len(state['deliveries']), 1)


if __name__ == '__main__':
    unittest.main()

