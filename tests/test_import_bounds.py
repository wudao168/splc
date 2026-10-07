import io
import unittest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from server.importer import metadata, parse_file, parse_paste, table_from_grid


class ImportBoundsTests(unittest.TestCase):
    def parse(self, book):
        out = io.BytesIO(); book.save(out)
        return parse_file(out.getvalue(), '询价.xlsx')

    def test_formatting_only_tail_does_not_inflate_limits(self):
        book = Workbook(); sheet = book.active
        sheet.append(['名称', '品牌或材质', '型号或图号', '技术说明', '数量'])
        for i in range(12):
            sheet.append(['料品'+str(i), '品牌', '型号', '说明', i+1])
        sheet.cell(100000, 256).fill = PatternFill('solid', fgColor='FFFF00')
        result = self.parse(book)
        self.assertEqual(len(result['tables'][0]['rows']), 12)
        self.assertEqual(result['tables'][0]['mapping']['spec'], 2)
        self.assertEqual(result['tables'][0]['mapping']['brand'], 1)
        self.assertEqual(result['tables'][0]['mapping']['description'], 3)
        self.assertFalse(any('超过' in w for w in result['warnings']))

    def test_repeated_headers_choose_populated_group(self):
        book = Workbook(); sheet = book.active
        sheet.append(['料号','名称','数量','料号','名称','数量'])
        sheet.append([None,None,None,'001','实际料品',3])
        result = self.parse(book)
        table = result['tables'][0]
        self.assertEqual(table['mapping']['name'], 4)
        self.assertEqual(len(table['rows']), 1)

    def test_real_content_over_limit_still_warns(self):
        book = Workbook(); sheet = book.active
        sheet.append(['名称','数量']); sheet.append(['料品',1])
        sheet.cell(2,151).value = '真实内容'
        result = self.parse(book)
        self.assertEqual(result['tables'], [])
        self.assertTrue(any('有效内容范围超过' in w for w in result['warnings']))

    def test_description_and_spec_synonyms(self):
        for description in ['描述', '说明', '技术说明']:
            for spec in ['规格', '型号', '规格型号', '型号或图号', '规格/型号']:
                with self.subTest(description=description, spec=spec):
                    table = parse_paste('名称\t'+spec+'\t'+description+'\t数量\n螺栓\tM20\t镀锌\t2', 'items')['tables'][0]
                    self.assertEqual(table['mapping']['spec'], 1)
                    self.assertEqual(table['mapping']['description'], 2)

    def test_pdf_quote_table_with_wrapped_cells(self):
        grid = [
            ['序号', '产品名称', '规格型号', '品牌', '数量', '单位', '单价\n（含税）'],
            ['1', '欧标电源线', '欧标三插头3*1.5平方\n字尾（全铜）1.5米', '国优', '10', '条', '17.50'],
        ]
        table = table_from_grid(grid, '第 1 页')
        self.assertEqual(table['mapping']['name'], 1)
        self.assertEqual(table['mapping']['price'], 6)
        self.assertEqual(table['rows'][0]['values'][2], '欧标三插头3*1.5平方\n字尾（全铜）1.5米')
        self.assertEqual(metadata('需    方: 上腾科技（广州）有限公司                     供    方: 供应商')['customer'],
                         '上腾科技（广州）有限公司')
