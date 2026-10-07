import io
import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from server.db import connect
from server import domain as dm
from server import lifecycle as lc
from server.importer import parse_file, parse_paste
from server.exports import delivery_data, xlsx_delivery, pdf_delivery, html_delivery


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.con = connect(self.temp.name)
        self.con.execute('INSERT INTO attachments VALUES(?,?,?,?)', ('a' * 64, 'invoice.pdf', 'application/pdf', dm.now()))
        self.oid = dm.create_order(self.con, {'customer':'测试客户', 'po':'PO-001', 'address':'广州测试路1号', 'contact':'采购员',
            'lines':[{'name':'化学螺栓','spec':'M20×260','brand':'测试品牌','description':'镀锌','quantity':30,'unit':'个','price':'10.01'}]})['id']
        self.lid = dm.row(self.con, 'SELECT id FROM order_lines WHERE order_id=?', (self.oid,))['id']
        dm.confirm_quote(self.con, self.oid)

    def tearDown(self):
        self.con.close()
        self.temp.cleanup()

    def purchase(self, n=30, amount='180', code='JD-001'):
        return dm.create_purchase(self.con, {'platform':'京东','shop':'测试店铺','platform_order':code,'amount':amount,
            'lines':[{'order_line_id':self.lid,'quantity':n,'purchase_quantity':n/10,'purchase_unit':'包','cost':amount}]})['id']

    def package(self, pid, quantity, tracking='TEST001'):
        pl = dm.row(self.con, 'SELECT id FROM purchase_lines WHERE purchase_id=?', (pid,))['id']
        return dm.create_package(self.con, {'carrier':'京东物流','tracking':tracking,'lines':[{'purchase_line_id':pl,'quantity':quantity}]})['id']

    def test_cancel_shared_purchase_preserves_other_order_and_history(self):
        second = dm.create_order(self.con, {'customer':'测试客户', 'po':'PO-002', 'address':'广州测试路1号',
            'lines':[{'name':'另一料品','quantity':2}]})['id']
        dm.confirm_quote(self.con, second)
        second_line = dm.row(self.con, 'SELECT id FROM order_lines WHERE order_id=?', (second,))['id']
        pid = dm.create_purchase(self.con, {'platform':'京东','shop':'测试店铺','platform_order':'SHARED-1','amount':30,
            'lines':[{'order_line_id':self.lid,'quantity':1,'cost':10},
                     {'order_line_id':second_line,'quantity':2,'cost':20}]})['id']
        first_purchase_line = dm.row(self.con, 'SELECT id FROM purchase_lines WHERE order_line_id=?', (self.lid,))['id']
        other_purchase_line = dm.row(self.con, 'SELECT id FROM purchase_lines WHERE order_line_id=?', (second_line,))['id']
        did = dm.create_deliveries(self.con, {'company':'测试公司','lines':[{'order_line_id':self.lid,'quantity':1}]})['ids'][0]
        self.con.execute("UPDATE deliveries SET shipped_at='' WHERE id=?", (did,))
        with self.assertRaisesRegex(ValueError, '送货单'):
            dm.cancel_order(self.con, self.oid, {'reason':'客户撤单'})
        dm.void_delivery(self.con, did)
        with self.assertRaisesRegex(ValueError, 'reason'):
            dm.cancel_order(self.con, self.oid, {})
        dm.cancel_order(self.con, self.oid, {'reason':'客户撤单'})
        state = dm.get_state(self.con)
        self.assertEqual(next(o for o in state['orders'] if o['id'] == self.oid)['status'], 'cancelled')
        self.assertEqual(next(o for o in state['orders'] if o['id'] == second)['status'], 'confirmed')
        self.assertEqual(len(state['purchases']), 1)
        self.assertEqual(len(state['deliveries']), 1)
        with self.assertRaisesRegex(ValueError, '没有可取消'):
            dm.cancel_order(self.con, self.oid, {'reason':'重复'})
        with self.assertRaisesRegex(ValueError, '已取消'):
            dm.create_package(self.con, {'carrier':'京东物流','tracking':'CANCELLED',
                'lines':[{'purchase_line_id':first_purchase_line,'quantity':1}]})
        with self.assertRaisesRegex(ValueError, '已取消'):
            dm.create_deliveries(self.con, {'company':'测试公司','lines':[{'order_line_id':self.lid,'quantity':1}]})
        self.assertTrue(dm.create_package(self.con, {'carrier':'京东物流','tracking':'ACTIVE',
            'lines':[{'purchase_line_id':other_purchase_line,'quantity':2}]})['id'])
        case = dm.get_state(self.con)['purchase_cases'][0]
        lc.update_purchase_case(self.con, case['id'], {'kind':'stock','location':'仓库 A','amount':0,'note':'改为库存'})
        lc.complete_purchase_case(self.con, case['id'])
        self.assertEqual(lc.purchase_balance(self.con, first_purchase_line)[0], 0)
        self.assertEqual(lc.purchase_balance(self.con, other_purchase_line)[0], 2)

    def test_quantity_and_duplicate_guards(self):
        pid = self.purchase(20, '120')
        with self.assertRaisesRegex(ValueError, '超过客户需求'):
            self.purchase(11, '60', 'JD-002')
        with self.assertRaisesRegex(ValueError, '已登记'):
            self.purchase(10, '60')
        self.package(pid, 12)
        with self.assertRaisesRegex(ValueError, '超过'):
            self.package(pid, 9, 'TEST002')
        self.package(pid, 8, 'TEST003')

    def test_money_and_quote_versions(self):
        self.assertEqual(dm.money('1.005'), 101)
        for invalid in ['nan','Infinity','-1']:
            with self.assertRaises(ValueError): dm.money(invalid)
        with self.assertRaisesRegex(ValueError, '已确认'):
            dm.update_quote(self.con, self.oid, {'prices':[{'id':self.lid,'price':11}]})
        dm.revise_quote(self.con, self.oid)
        dm.update_quote(self.con, self.oid, {'prices':[{'id':self.lid,'price':11}]})
        dm.confirm_quote(self.con, self.oid)
        snapshots = dm.rows(self.con, 'SELECT * FROM quotes ORDER BY version')
        self.assertEqual(len(snapshots),2)
        self.assertEqual(json.loads(snapshots[0]['snapshot'])[0]['price_cents'],1001)
        self.assertEqual(json.loads(snapshots[1]['snapshot'])[0]['price_cents'],1100)

    def test_draft_edit_before_purchase_only(self):
        dm.revise_quote(self.con,self.oid)
        dm.edit_order(self.con,self.oid,{'customer':'测试客户','po':'PO-001','address':'更正地址','lines':[{'name':'新料品','quantity':2}]})
        self.assertEqual(dm.row(self.con,'SELECT address FROM orders WHERE id=?',(self.oid,))['address'],'更正地址')
        self.lid = dm.row(self.con,'SELECT id FROM order_lines WHERE order_id=?',(self.oid,))['id']
        dm.confirm_quote(self.con,self.oid)
        self.purchase(1,'1')
        dm.revise_quote(self.con,self.oid)
        with self.assertRaisesRegex(ValueError,'关联采购'):
            dm.edit_order(self.con,self.oid,{'customer':'测试客户','po':'PO-001','address':'错误地址','lines':[{'name':'新料品','quantity':3}]})

    def test_purchase_cost_may_differ_from_payment(self):
        pid = dm.create_purchase(self.con, {'platform':'淘宝','shop':'店铺','platform_order':'T1','amount':100,
            'lines':[{'order_line_id':self.lid,'quantity':1,'cost':99}]})['id']
        purchase = dm.row(self.con, 'SELECT amount_cents FROM purchases WHERE id=?', (pid,))
        line = dm.row(self.con, 'SELECT cost_cents FROM purchase_lines WHERE purchase_id=?', (pid,))
        self.assertEqual((purchase['amount_cents'], line['cost_cents']), (10000, 9900))

    def test_purchase_attachments_can_be_added_and_removed(self):
        second = 'b' * 64
        self.con.execute('INSERT INTO attachments VALUES(?,?,?,?)', (second, 'receipt.png', 'image/png', dm.now()))
        pid = dm.create_purchase(self.con, {'platform':'京东','shop':'店铺','platform_order':'FILES-1','amount':10,
            'attachment_ids':['a' * 64, second], 'lines':[{'order_line_id':self.lid,'quantity':1}]})['id']
        self.assertEqual([file['name'] for file in dm.get_state(self.con)['purchases'][0]['attachments']], ['invoice.pdf', 'receipt.png'])
        dm.remove_purchase_attachment(self.con, pid, {'attachment_id':'a' * 64})
        self.assertEqual(dm.row(self.con, 'SELECT source_id FROM purchases WHERE id=?', (pid,))['source_id'], second)
        dm.add_purchase_attachments(self.con, pid, {'attachment_ids':['a' * 64]})
        self.assertEqual(len(dm.get_state(self.con)['purchases'][0]['attachments']), 2)
        with self.assertRaisesRegex(ValueError, '附件不存在'):
            dm.add_purchase_attachments(self.con, pid, {'attachment_ids':['c' * 64]})

    def test_existing_purchase_source_is_migrated_to_attachment_list(self):
        pid = self.purchase(1, '10', 'FILES-OLD')
        self.con.execute('UPDATE purchases SET source_id=? WHERE id=?', ('a' * 64, pid))
        self.con.commit()
        other = connect(self.temp.name)
        try:
            self.assertEqual([file['name'] for file in dm.get_state(other)['purchases'][0]['attachments']], ['invoice.pdf'])
        finally:
            other.close()

    def test_split_packages_invoice_clock_and_partial_invoice(self):
        pid = self.purchase()
        k1 = self.package(pid, 20)
        k2 = self.package(pid, 10,'TEST002')
        signed = (date.fromisoformat(dm.today()) - timedelta(days=7)).isoformat()
        dm.update_tracking(self.con,k1,{'status':'已签收','signed_at':signed})
        self.assertFalse(dm.get_state(self.con)['purchases'][0]['overdue'])
        dm.update_tracking(self.con,k2,{'status':'已签收','signed_at':signed})
        p = dm.get_state(self.con)['purchases'][0]
        self.assertEqual(p['effective_due'],dm.today())
        self.assertTrue(p['overdue'])
        invoice = {'number':'INV001','seller':'测试公司','issued_date':dm.today(),'amount':100,'attachment_id':'a'*64,
                   'allocations':[{'purchase_id':pid,'amount':100}]}
        dm.create_invoice(self.con, invoice)
        p = dm.get_state(self.con)['purchases'][0]
        self.assertEqual(p['remaining_cents'],8000)
        self.assertTrue(p['overdue'])
        with self.assertRaisesRegex(ValueError,'已登记'):
            dm.create_invoice(self.con,invoice)
        invoice.update(number='INV002',amount=80,allocations=[{'purchase_id':pid,'amount':80}])
        dm.create_invoice(self.con,invoice)
        p = dm.get_state(self.con)['purchases'][0]
        self.assertEqual(p['receipt_status'],'已收齐')
        self.assertFalse(p['overdue'])

    def test_unallocated_quantity_prevents_invoice_clock(self):
        pid = self.purchase()
        kid = self.package(pid, 20)
        dm.update_tracking(self.con,kid,{'status':'已签收','signed_at':'2026-01-01'})
        self.assertEqual(dm.get_state(self.con)['purchases'][0]['effective_due'],'')

    def test_followup_date_and_shipped_quantities(self):
        pid = self.purchase()
        kid = self.package(pid, 20)
        dm.update_tracking(self.con,kid,{'status':'运输中'})
        dm.follow_invoice(self.con,pid,{'invoice_stage':'已申请待开票','next_followup':dm.today(),'followup':'待店铺回复'})
        state = dm.get_state(self.con)
        self.assertTrue(state['purchases'][0]['followup_due'])
        self.assertEqual(state['order_lines'][0]['shipped'],20)
        self.assertEqual(state['order_lines'][0]['signed'],0)
        dm.update_tracking(self.con,kid,{'status':'已签收','signed_at':dm.today()})
        self.assertEqual(dm.get_state(self.con)['order_lines'][0]['signed'],20)

    def test_delivery_snapshot_export_void_and_reprint(self):
        pid = self.purchase()
        kid = self.package(pid,30)
        klid = dm.row(self.con,'SELECT id FROM package_lines WHERE package_id=?',(kid,))['id']
        request = {'company':'测试供货公司','mode':'combined','lines':[{'package_line_id':klid,'quantity':20}]}
        with self.assertRaisesRegex(ValueError, '待揽收'):
            dm.create_deliveries(self.con,request)
        dm.update_tracking(self.con,kid,{'status':'运输中'})
        did = dm.create_deliveries(self.con,request)['ids'][0]
        self.assertEqual(dm.row(self.con, 'SELECT number FROM deliveries WHERE id=?', (did,))['number'], 'SHD-PO-001-00001')
        with self.assertRaisesRegex(ValueError, '超过'):
            dm.create_deliveries(self.con,request)
        document = delivery_data(self.con,did)
        self.con.execute('UPDATE order_lines SET name=? WHERE id=?',('已修改名称',self.lid))
        self.assertEqual(delivery_data(self.con,did)['lines'][0]['name'],'化学螺栓')
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(xlsx_delivery(document)))
        contents = ' '.join(str(c.value) for row in wb.active for c in row if c.value is not None)
        for excluded in ['单价','金额','客户签收']:
            self.assertNotIn(excluded,contents)
            self.assertNotIn(excluded,html_delivery(document).decode())
        self.assertIn('化学螺栓',contents)
        self.assertEqual(wb.active['D7'].value,20)
        self.assertEqual(wb.active['B7'].alignment.horizontal, 'center')
        self.assertEqual(wb.active['G7'].alignment.vertical, 'center')
        self.assertIn('客户 PO：PO-001', wb.active['A1'].value)
        self.assertEqual(wb.active.page_setup.orientation, 'portrait')
        self.assertEqual(wb.active['G6'].value, '项目号')
        self.assertEqual(wb.active['H6'].value, '子项目号')
        self.assertEqual(wb.active['F4'].alignment.horizontal, 'center')
        self.assertIn('收货地址：广州测试路1号', contents)
        self.assertIn('text-align:center;vertical-align:middle', html_delivery(document).decode())
        self.assertTrue(pdf_delivery(document).startswith(b'%PDF'))
        self.assertEqual(dm.get_state(self.con)['package_lines'][0]['delivered'],20)
        self.assertTrue(dm.row(self.con, 'SELECT shipped_at FROM deliveries WHERE id=?', (did,))['shipped_at'])
        self.con.execute("UPDATE deliveries SET shipped_at='' WHERE id=?", (did,))
        dm.void_delivery(self.con,did)
        self.assertEqual(dm.get_state(self.con)['package_lines'][0]['delivered'],0)
        dm.create_deliveries(self.con,request)

    def test_delivery_by_order_line_without_tracking_or_shipping_status(self):
        request = {'company':'测试供货公司', 'lines':[{'order_line_id':self.lid,'quantity':10}]}
        did = dm.create_deliveries(self.con, request)['ids'][0]
        self.assertEqual(dm.row(self.con, 'SELECT number FROM deliveries WHERE id=?', (did,))['number'], 'SHD-PO-001-00001')
        line = delivery_data(self.con, did)['lines'][0]
        self.assertEqual((line['tracking'], line['quantity']), ('', 10))
        pid = self.purchase()
        self.package(pid, 20)
        did2 = dm.create_deliveries(self.con, {'company':'测试供货公司', 'lines':[{'order_line_id':self.lid,'quantity':20}]})['ids'][0]
        self.assertEqual(dm.row(self.con, 'SELECT number FROM deliveries WHERE id=?', (did2,))['number'], 'SHD-PO-001-00002')
        self.assertEqual(delivery_data(self.con, did2)['lines'][0]['tracking'], 'TEST001')
        with self.assertRaisesRegex(ValueError, '超过订单剩余数量'):
            dm.create_deliveries(self.con, request)
        self.con.execute("UPDATE deliveries SET shipped_at='' WHERE id=?", (did,))
        dm.void_delivery(self.con, did)
        did3 = dm.create_deliveries(self.con, request)['ids'][0]
        self.assertEqual(did3, did2 + 1)
        self.assertEqual(dm.row(self.con, 'SELECT number FROM deliveries WHERE id=?', (did3,))['number'], 'SHD-PO-001-00003')

    def test_missing_attachment_and_no_invoice_reason(self):
        pid = self.purchase()
        with self.assertRaises(ValueError):
            dm.create_invoice(self.con,{'number':'I','seller':'S','issued_date':dm.today(),'amount':1,'allocations':[{'purchase_id':pid,'amount':1}]})
        with self.assertRaisesRegex(ValueError,'原因'):
            dm.follow_invoice(self.con,pid,{'invoice_stage':'不需开票'})

    def test_import_xlsx_and_paste(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(['客户名称：测试公司'])
        ws.append(['品名','型号','品牌','数量','单位'])
        ws.append(['密封胶','001-A',None,10,'支'])
        ws.append(['合计',None,None,10])
        out = io.BytesIO(); wb.save(out)
        result = parse_file(out.getvalue(),'客户.xlsx')
        self.assertEqual(result['metadata']['customer'],'测试公司')
        self.assertEqual(len(result['tables'][0]['rows']),1)
        self.assertEqual(result['tables'][0]['rows'][0]['values'][1],'001-A')
        parsed = parse_paste('料品名称\t料品规格\t数量\n螺栓\tM20\t30','items')
        self.assertEqual(parsed['tables'][0]['mapping']['quantity'],2)
        result = parse_paste('订单编号：12345678901234567890\n店铺：测试店铺\n实付款：￥1,234.56\n下单时间：2026/10/3 12:30','purchase')
        self.assertEqual(result['amount'],'1234.56')
        self.assertEqual(result['platform_order'],'12345678901234567890')
        self.assertEqual(result['shop'],'测试店铺')
        self.assertEqual(result['purchased_date'],'2026-10-03')
        screenshot_text = '商品 总 价 ¥81.20\n实 付款 ~ ¥65.00\n订单 编号 2234567890123456789 复制\n创建 时 间 2026-10-02 21:22:07'
        extracted = parse_paste(screenshot_text, 'purchase')
        self.assertEqual(extracted['platform_order'], '2234567890123456789')
        self.assertEqual(extracted['amount'], '65.00')
        self.assertEqual(extracted['purchased_date'], '2026-10-02')
        shipment = parse_paste('卖家：淘宝店铺\n物流公司：中通快递\n运单号码：ZT123456789', 'purchase')
        self.assertEqual(shipment['shop'],'淘宝店铺')
        self.assertEqual(shipment['carrier'],'中通快递')
        self.assertEqual(shipment['tracking'],'ZT123456789')
        taobao_text = ('包裹1(共1件)\n极兔速递\nJT1234567890123\n展开全部商品\n'
                       '测试服饰店\n查看交易快照\n商品总价\n￥81.20\n'
                       '实付款\n￥65.00\n订单编号\n2234567890123456789\n复制\n'
                       '创建时间\n2026-10-02 21:22:07')
        extracted = parse_paste(taobao_text, 'purchase')
        self.assertEqual(extracted, {'platform_order':'2234567890123456789', 'shop':'测试服饰店',
                                     'amount':'65.00', 'purchased_date':'2026-10-02',
                                     'carrier':'极兔速递', 'tracking':'JT1234567890123'})
        with self.assertRaisesRegex(ValueError, '未识别到订单号'):
            parse_paste('没有订单信息', 'purchase')

    def test_merge_different_customers_rejected(self):
        pid = self.purchase()
        kid = self.package(pid,30)
        dm.update_tracking(self.con,kid,{'status':'运输中'})
        oid = dm.create_order(self.con,{'customer':'另一客户','po':'PO-002','address':'深圳','lines':[{'name':'阀门','quantity':1}]})['id']
        lid = dm.row(self.con,'SELECT id FROM order_lines WHERE order_id=?',(oid,))['id']
        dm.confirm_quote(self.con,oid)
        p2 = dm.create_purchase(self.con,{'platform':'淘宝','shop':'A','platform_order':'T2','amount':1,'lines':[{'order_line_id':lid,'quantity':1,'cost':1}]})['id']
        k2 = self.package(p2,1,'TEST003')
        dm.update_tracking(self.con,k2,{'status':'运输中'})
        kls = dm.rows(self.con,'SELECT id,quantity FROM package_lines')
        d = {'company':'公司','lines':[{'package_line_id':x['id'],'quantity':x['quantity']} for x in kls]}
        with self.assertRaisesRegex(ValueError,'合并开单'):
            dm.create_deliveries(self.con,d)
        d['mode'] = 'separate'
        self.assertEqual(len(dm.create_deliveries(self.con,d)['ids']),2)


if __name__ == '__main__':
    unittest.main()
